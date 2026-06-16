import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Substation Workflows"

# ── Data: (module, sub_module, url1, url2) ────────────────────────────────────
# Broken URL replacements:
#   lauwtjunnji.weebly.com  → https://www.academia.edu/8347063/FIDIC_Conditions_of_Contract_for_EPC_Turnkey_Projects_doc
#   pmri.in/epccompletion/  → https://www.projectmanager.com/blog/managing-construction-projects
#   petroedgeasia.net       → https://www.lexology.com/library/detail.aspx?g=2de63e41-6eef-432b-a5c8-0026bbeac029
#   electra.cigre.org       → https://www.electrical4u.com/electrical-power-substation-engineering-and-layout/
#   mayerbrown.com PDF      → https://projectmanagementformula.com/engineering-procurement-and-construction-epc-contract/
#   pillsburylaw.com PDF    → https://constructionfront.com/epc-contract/

FIDIC_URL  = "https://www.academia.edu/8347063/FIDIC_Conditions_of_Contract_for_EPC_Turnkey_Projects_doc"
PMRI_URL   = "https://www.projectmanager.com/blog/managing-construction-projects"
PETRO_URL  = "https://www.lexology.com/library/detail.aspx?g=2de63e41-6eef-432b-a5c8-0026bbeac029"
CIGRE_URL  = "https://www.electrical4u.com/electrical-power-substation-engineering-and-layout/"
MAYER_URL  = "https://projectmanagementformula.com/engineering-procurement-and-construction-epc-contract/"
PILLS_URL  = "https://constructionfront.com/epc-contract/"

data = [
    # WORKFLOW 1 — Bid Process
    ("Bid Process in Substation EPC Projects",
     "Tender / NIT Received – Opportunity Identification",
     "https://constructionfront.com/epc-contract/",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents"),

    ("",
     "Bid / No-Bid Decision – Opportunity Evaluation",
     "https://www.procore.com/library/epc-contractors",
     "https://prismecs.com/blog/understanding-the-lifecycle-of-epc-solutions-services"),

    ("",
     "Eligibility & Pre-Bid – Qualification Verification",
     "https://www.euci.com/event_post/epc-contracts/",
     "https://dancumberlandlabs.com/blog/epc-engineering-procurement-construction/"),

    ("",
     "Scope Review & Quantity Take-Off",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process"),

    ("",
     "Vendor & Subcontractor Quotations – Market Cost Collection",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://electrical-engineering-portal.com/checks-for-successful-substation-factory-acceptance-testing"),

    ("",
     "Parallel Cost Estimation – Detailed Cost Calculation",
     "https://rkstrainings.com/what-is-project-management-in-epc/",
     "https://prismecs.com/blog/understanding-the-lifecycle-of-epc-solutions-services"),

    ("",
     "Base Cost Estimation – Total Project Cost Consolidation",
     "https://rkstrainings.com/what-is-project-management-in-epc/",
     "https://www.pmi.org/learning/library/realizing-engineering-procurement-construction-projects-7173"),

    ("",
     "Risk & Contingency Analysis – Project Risk Assessment",
     MAYER_URL,
     "https://www.pmi.org/learning/library/realizing-engineering-procurement-construction-projects-7173"),

    ("",
     "Margin Finalization – Profit Strategy Definition",
     "https://www.euci.com/event_post/epc-contracts/",
     PILLS_URL),

    ("",
     "Technical & Commercial Bid Preparation – Proposal Creation",
     PILLS_URL,
     "https://fidic.org/sites/default/files/The%20FIDIC%20Contracts%20Guide.pdf"),

    ("",
     "Client Evaluation – Tender Assessment by Client",
     "https://www.procore.com/library/epc-contractors",
     "https://www.sora.ro/2026/05/31/epc-versus-fidic-contracts/"),

    ("",
     "Bid Award (LOA) – Project Successfully Won",
     "https://relgrow.com/resources/dlp-full-form-in-construction/",
     PILLS_URL),

    ("",
     "Contract Finalization & PBG – Legal Agreement Completion",
     "https://fidic.org/sites/default/files/The%20FIDIC%20Contracts%20Guide.pdf",
     FIDIC_URL),

    # WORKFLOW 2 — Contract Finalization
    ("Contract Finalization and Project Mobilization Process",
     "LOA / LOI Received – Project Award Confirmation",
     PMRI_URL,
     "https://relgrow.com/resources/dlp-full-form-in-construction/"),

    ("",
     "Draft and Review Contract – Contract Preparation",
     "https://fidic.org/sites/default/files/The%20FIDIC%20Contracts%20Guide.pdf",
     PILLS_URL),

    ("",
     "Negotiation and Finalization – Commercial and Legal Alignment",
     PETRO_URL,
     "https://fidic.org/sites/default/files/The%20FIDIC%20Contracts%20Guide.pdf"),

    ("",
     "Contract Signing – Legal Agreement Execution",
     "https://fidic.org/sites/default/files/The%20FIDIC%20Contracts%20Guide.pdf",
     "https://constructionfront.com/epc-contract/"),

    ("",
     "Parallel Compliance Activities – Financial and Regulatory Commitments",
     FIDIC_URL,
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/"),

    ("",
     "Contract Effective – Contract Activation",
     PMRI_URL,
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    ("",
     "Bid to Execution Handover – Ownership Transfer",
     "https://projectmanagement123.com/project-completion-and-handover-procedure/",
     PMRI_URL),

    ("",
     "Project Planning Activities – Execution Preparation",
     "https://www.pmi.org/learning/library/realizing-engineering-procurement-construction-projects-7173",
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    # WORKFLOW 3 — Project Mobilization
    ("Project Mobilization and Site Establishment Process",
     "Mobilization Advance and Bank Guarantee Processing",
     FIDIC_URL,
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/"),

    ("",
     "Site Possession – Physical Site Handover",
     "https://www.substationfaults.com/electrical-substation-construction/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Resource Mobilization – Manpower and Equipment Deployment",
     "https://prismecs.com/blog/understanding-the-lifecycle-of-epc-solutions-services",
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    ("",
     "Temporary Infrastructure and Site Office Setup",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     "https://www.substationfaults.com/electrical-substation-construction/"),

    ("",
     "HSE Plan and Safety Induction Setup",
     CIGRE_URL,
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Site Established – Ready for Execution",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     "https://www.substationfaults.com/electrical-substation-construction/"),

    # WORKFLOW 4 — Survey & Geotechnical
    ("Survey, Geotechnical Investigation and Design Release Process",
     "Benchmark Points and GPS Grid Layout",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process"),

    ("",
     "Topographic Survey",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process",
     "https://www.substationfaults.com/electrical-substation-construction/"),

    ("",
     "Soil Investigation and Geotechnical Testing",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process"),

    ("",
     "Geotechnical Report Preparation",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process",
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    ("",
     "Technical Review and Client Approval of Reports",
     "https://electrical-engineering-portal.com/turnkey-substation-technical-documentation",
     "https://www.pmi.org/learning/library/realizing-engineering-procurement-construction-projects-7173"),

    ("",
     "Design Basis Finalized – Engineering Inputs Released",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://electrical-engineering-portal.com/turnkey-substation-technical-documentation"),

    # WORKFLOW 5 — Detailed Engineering
    ("Detailed Engineering and IFC Release Process",
     "Design Basis and Engineering Inputs Collection",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://www.academia.edu/93132452/ElECTRIC_POWER_SUBSTATIONS_ENGINEERING"),

    ("",
     "Single Line Diagram (SLD) Preparation",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://keentelengineering.com/ieee-compliant-ehv-hv-mv-substation-design-services-by-keentel-engineering"),

    ("",
     "General Arrangement and Layout Design",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process",
     "https://www.academia.edu/93132452/ElECTRIC_POWER_SUBSTATIONS_ENGINEERING"),

    ("",
     "Civil and Structural Design",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process"),

    ("",
     "Protection, Control and SCADA Design",
     "https://seclab.illinois.edu/wp-content/uploads/2011/03/iec61850-intro.pdf",
     "https://library.e.abb.com/public/5ea2620c3dec4376a1540fc2b07a18b7/IEC%2061850%20Overview_Self.pdf"),

    ("",
     "Grounding and Earthing Design",
     "https://keentelengineering.com/ieee-compliant-ehv-hv-mv-substation-design-services-by-keentel-engineering",
     "https://www.academia.edu/93132452/ElECTRIC_POWER_SUBSTATIONS_ENGINEERING"),

    ("",
     "Cable and Raceway Design",
     "https://electrical-engineering-portal.com/power-substation-design-engineering",
     "https://keentelengineering.com/ieee-compliant-ehv-hv-mv-substation-design-services-by-keentel-engineering"),

    ("",
     "Engineering Calculations – Fault, Battery, Lightning, Grounding",
     "https://keentelengineering.com/ieee-compliant-ehv-hv-mv-substation-design-services-by-keentel-engineering",
     "https://www.academia.edu/93132452/ElECTRIC_POWER_SUBSTATIONS_ENGINEERING"),

    # WORKFLOW 6 — Submit for Approval
    ("Submit for Approval – Client Submission Process",
     "Drawing and Document Submission to Client",
     "https://electrical-engineering-portal.com/turnkey-substation-technical-documentation",
     "https://keentelengineering.com/a-guide-to-the-substation-design-process"),

    ("",
     "Client Technical Review and Comments",
     "https://electrical-engineering-portal.com/turnkey-substation-technical-documentation",
     "https://www.pmi.org/learning/library/realizing-engineering-procurement-construction-projects-7173"),

    ("",
     "Revision Incorporation and Resubmission",
     "https://electrical-engineering-portal.com/turnkey-substation-technical-documentation",
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    ("",
     "Approved – IFC (Issued for Construction) Release",
     "https://electrical-engineering-portal.com/turnkey-substation-technical-documentation",
     "https://www.academia.edu/93132452/ElECTRIC_POWER_SUBSTATIONS_ENGINEERING"),

    # WORKFLOW 7 — Procurement
    ("Procurement, Vendor Approval and Manufacturing Process",
     "Procurement Plan Preparation",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    ("",
     "Request for Quotation (RFQ) and Vendor Identification",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://electrical-engineering-portal.com/checks-for-successful-substation-factory-acceptance-testing"),

    ("",
     "Bid Evaluation and Techno-Commercial Comparison",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://www.euci.com/event_post/epc-contracts/"),

    ("",
     "Purchase Order (PO) Issuance and Vendor Approval",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://keentelengineering.com/substation-testing-lifecycle-fat-vs-sat"),

    ("",
     "Advance Payment and Bank Guarantee Processing",
     FIDIC_URL,
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/"),

    # WORKFLOW 8 — Manufacturing & FAT
    ("Manufacturing Execution, FAT Testing and Dispatch Process",
     "Manufacturing Kickoff – Production Planning",
     "https://electrical-engineering-portal.com/checks-for-successful-substation-factory-acceptance-testing",
     "https://keentelengineering.com/substation-testing-lifecycle-fat-vs-sat"),

    ("",
     "In Production – Equipment Manufacturing Process",
     "https://electrical-engineering-portal.com/substation-equipment-type-testing-factory-site-acceptance-testing-fat-sat",
     "https://electrical-engineering-portal.com/transformer-factory-acceptance-tests-final-inspections"),

    ("",
     "Stage Inspection – In-Process Quality Verification",
     "https://electrical-engineering-portal.com/substation-equipment-type-testing-factory-site-acceptance-testing-fat-sat",
     "https://keentelengineering.com/substation-testing-lifecycle-fat-vs-sat"),

    ("",
     "Factory Acceptance Testing (FAT)",
     "https://electrical-engineering-portal.com/checks-for-successful-substation-factory-acceptance-testing",
     "https://electrical-engineering-portal.com/res3/FAT-SAT-OF-electrical-and-automation-systems.pdf"),

    ("",
     "Dispatch Clearance and Packing",
     "https://keentelengineering.com/substation-testing-lifecycle-fat-vs-sat",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents"),

    # WORKFLOW 9 — Logistics
    ("Logistics, Transportation and Material Receipt Process",
     "Transport Planning and Route Survey",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    ("",
     "Equipment Dispatch from Vendor Facility",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://keentelengineering.com/substation-testing-lifecycle-fat-vs-sat"),

    ("",
     "In-Transit Monitoring and Tracking",
     "https://rkstrainings.com/what-is-project-management-in-epc/",
     "https://prismecs.com/blog/understanding-the-lifecycle-of-epc-solutions-services"),

    ("",
     "Site Unloading and Placement",
     "https://www.substationfaults.com/electrical-substation-construction/",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning"),

    # WORKFLOW 10 — Material Receipt
    ("Material Receipt, Inspection and Installation Release Process",
     "Site Preparation for Material Receipt",
     "https://www.substationfaults.com/electrical-substation-construction/",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning"),

    ("",
     "Material Unloading and Receipt",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://www.substationfaults.com/electrical-substation-construction/"),

    ("",
     "Incoming Inspection and Quality Verification",
     "https://electrical-engineering-portal.com/substation-equipment-type-testing-factory-site-acceptance-testing-fat-sat",
     "https://keentelengineering.com/substation-testing-lifecycle-fat-vs-sat"),

    ("",
     "SDR / Claim Management – Discrepancy Resolution Process",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/"),

    ("",
     "Storage and Inventory Management",
     "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents",
     "https://rkstrainings.com/what-is-project-management-in-epc/"),

    ("",
     "Released for Installation",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     "https://www.substationfaults.com/electrical-substation-construction/"),

    # WORKFLOW 11 — Civil Construction
    ("Civil Construction and Structural Readiness Process",
     "Site Clearing, Grading and Levelling",
     "https://www.substationfaults.com/electrical-substation-construction/",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning"),

    ("",
     "Below Grade Works – Cable Trenches, Conduits, Ground Grid",
     "https://www.substationfaults.com/electrical-substation-construction/",
     CIGRE_URL),

    ("",
     "Foundation Construction",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     "https://www.substationfaults.com/electrical-substation-construction/"),

    ("",
     "Control Room and Buildings Construction",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     CIGRE_URL),

    ("",
     "Drainage, Fencing and Stone Verge",
     "https://www.substationfaults.com/electrical-substation-construction/",
     CIGRE_URL),

    ("",
     "Civil Readiness Check – Structural Work Complete",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning"),

    # WORKFLOW 12 — Equipment Erection
    ("Equipment Erection, Mechanical Installation and Cabling Readiness Process",
     "Prerequisite Verification – Civil and Materials Ready",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/"),

    ("",
     "Structural Steel Erection",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     CIGRE_URL),

    ("",
     "Power Transformer and Major Equipment Installation",
     "https://electrical-engineering-portal.com/transformer-factory-acceptance-tests-final-inspections",
     "https://www.substationfaults.com/electrical-substation-construction/"),

    ("",
     "Bus Bar and Conductor Installation",
     "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning",
     "https://keentelengineering.com/ieee-compliant-ehv-hv-mv-substation-design-services-by-keentel-engineering"),

    ("",
     "Mechanical Completion Check",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    # WORKFLOW 13 — Cabling
    ("Cabling, Termination, Testing and SAS Readiness Process",
     "Power and Control Cable Laying",
     "https://www.substationfaults.com/electrical-substation-construction/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Cable Termination at Equipment and Panels",
     "https://www.substationfaults.com/electrical-substation-construction/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Point-to-Point Continuity Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    ("",
     "SAS Panel Installation and IED Configuration",
     "https://instrumentationtools.com/scada-for-substation-automation/",
     "https://seclab.illinois.edu/wp-content/uploads/2011/03/iec61850-intro.pdf"),

    ("",
     "SAS Readiness Check",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-4-scada-and-comms-systems",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    # WORKFLOW 14 — SAS / SCADA Integration
    ("SAS / SCADA Integration, Functional Testing and Pre-Commissioning Readiness Process",
     "IED Configuration and Protection Logic Programming",
     "https://seclab.illinois.edu/wp-content/uploads/2011/03/iec61850-intro.pdf",
     "https://library.e.abb.com/public/5ea2620c3dec4376a1540fc2b07a18b7/IEC%2061850%20Overview_Self.pdf"),

    ("",
     "SCADA Integration and Communication Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-4-scada-and-comms-systems",
     "https://instrumentationtools.com/scada-for-substation-automation/"),

    ("",
     "Functional Testing of Protection Schemes",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/"),

    ("",
     "HMI and Local Control Panel Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-4-scada-and-comms-systems",
     "https://mbcontrol.com/understanding-substation-automation-systems/"),

    ("",
     "Pre-Commissioning Readiness Declaration",
     "https://powersynchro.com/substation-commissioning-checklist/",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/"),

    # WORKFLOW 15 — Pre-Commissioning
    ("Pre-Commissioning Testing, System Validation and Commissioning Readiness Process",
     "Insulation Resistance (IR) Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "CT / PT Ratio, Polarity and Burden Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Protection Relay Secondary Injection Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/",
     "https://elliotengineeringinc.com/testing-and-commissioning/"),

    ("",
     "Circuit Breaker Timing and Mechanical Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    ("",
     "Transformer Differential and REF Stability Testing",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/"),

    ("",
     "Ground Grid Continuity and Earthing Verification",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Punch List Identification and Rectification",
     "https://powersynchro.com/substation-commissioning-checklist/",
     PMRI_URL),

    ("",
     "Commissioning Readiness Certificate Issued",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    # WORKFLOW 16 — Final Commissioning
    ("Final Commissioning, Initial Charging and Energization Readiness Process",
     "Pre-Commissioning Completion Verification",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    ("",
     "External Clearances – Grid / SLDC / Regulatory Approvals",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Parallel External Clearance Activities – Operational Authorization Process",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    ("",
     "Initial Charging – Step-by-Step Energization",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Protection Selectivity Verification Under Live Conditions",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/",
     "https://elliotengineeringinc.com/testing-and-commissioning/"),

    ("",
     "Energization Readiness Confirmed",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    # WORKFLOW 17 — Final Energization
    ("Final Energization and Grid Synchronization Process",
     "Grid Authorization Received",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Charge and Synchronize – Permanent Grid Connection Process",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    ("",
     "Load Flow Verification",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Metering and Billing Meter Verification",
     "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/",
     "https://powersynchro.com/substation-commissioning-checklist/"),

    # WORKFLOW 18 — Trial Run
    ("Trial Run and Performance Stabilization Process",
     "Trial Operation Period",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"),

    ("",
     "Performance Monitoring – Live Data Collection",
     "https://www.ncbi.nlm.nih.gov/pmc/articles/PMC12431182/",
     "https://instrumentationtools.com/scada-for-substation-automation/"),

    ("",
     "Defect Identification and Correction During Trial Run",
     "https://relgrow.com/resources/dlp-full-form-in-construction/",
     "https://www.mastt.com/blogs/defects-liability-period"),

    ("",
     "Performance Stabilization Confirmed",
     "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/",
     PMRI_URL),

    # WORKFLOW 19 — Project Closure
    ("Project Closure, Client Acceptance and Final Handover Process",
     "Punch List Final Completion",
     PMRI_URL,
     "https://projectmanagement123.com/project-completion-and-handover-procedure/"),

    ("",
     "Client Walk-through and Acceptance Inspection",
     PMRI_URL,
     "https://projectmanagement123.com/project-completion-and-handover-procedure/"),

    ("",
     "As-Built Documentation Handover",
     "https://projectmanagement123.com/project-completion-and-handover-procedure/",
     PMRI_URL),

    ("",
     "Certificate of Completion / Provisional Acceptance Certificate (PAC)",
     "https://relgrow.com/resources/dlp-full-form-in-construction/",
     "https://www.mastt.com/blogs/defects-liability-period"),

    ("",
     "Contractor Closeout Submittals",
     "https://projectmanagement123.com/project-completion-and-handover-procedure/",
     PMRI_URL),

    # WORKFLOW 20 — DLP
    ("Defect Liability Period (DLP), Warranty Support and Final Acceptance Process",
     "DLP Period Commencement",
     "https://www.mastt.com/blogs/defects-liability-period",
     "https://www.designingbuildings.co.uk/wiki/Defects_liability_period_DLP"),

    ("",
     "Defect Attendance – Identification and Repair",
     "https://www.mastt.com/blogs/defects-liability-period",
     "https://www.turtons.com/blog/what-is-the-defects-liability-period"),

    ("",
     "Recurring Defect Investigation and Permanent Resolution",
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/",
     "https://www.mastt.com/blogs/defects-liability-period"),

    ("",
     "Warranty Claims Management",
     "https://www.mastt.com/blogs/defects-liability-period",
     FIDIC_URL),

    # WORKFLOW 21 — FAC & Closeout
    ("Final Acceptance Certificate (FAC), Contract Closure and Project Closeout Process",
     "All DLP Defects Rectified – Final Inspection Request",
     "https://relgrow.com/resources/dlp-full-form-in-construction/",
     PMRI_URL),

    ("",
     "Final Acceptance Certificate (FAC) Issuance",
     "https://relgrow.com/resources/dlp-full-form-in-construction/",
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/"),

    ("",
     "Final Payment and Retention Release",
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/",
     FIDIC_URL),

    ("",
     "Contract Closure and Project Closeout",
     "https://projectmanagement123.com/project-completion-and-handover-procedure/",
     "https://www.pmi.org/learning/library/realizing-engineering-procurement-construction-projects-7173"),

    ("",
     "Rework Loop – Unresolved Issues Resolution",
     "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/",
     "https://www.mastt.com/blogs/defects-liability-period"),
]

# ── Build rows: each (module, sub_module, url1, url2) → two Excel rows ─────────
# Row A: module_name (bold if first of workflow), sub_module, url1
# Row B: "",          "",          url2
excel_rows = []
for (mod, sub, url1, url2) in data:
    excel_rows.append((mod, sub, url1))
    excel_rows.append(("", "", url2))

# ── Build Workbook ─────────────────────────────────────────────────────────────
headers = ["Module", "Sub-Module / Step", "Source URL"]

thin = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

hdr_font  = Font(bold=True, size=11)
hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = hdr_font
    cell.alignment = hdr_align
    cell.border = border

wrap_top  = Alignment(vertical='top', wrap_text=True)
link_font = Font(color="0563C1", underline="single")

for row_num, (mod, sub, url) in enumerate(excel_rows, 2):
    # Col A — Module
    c1 = ws.cell(row=row_num, column=1, value=mod)
    c1.alignment = wrap_top
    c1.border = border
    if mod:
        c1.font = Font(bold=True)

    # Col B — Sub-Module
    c2 = ws.cell(row=row_num, column=2, value=sub)
    c2.alignment = wrap_top
    c2.border = border

    # Col C — Source URL (clickable)
    c3 = ws.cell(row=row_num, column=3, value=url)
    if url:
        c3.hyperlink = url
        c3.font = Font(color="0563C1", underline="single")
    c3.alignment = wrap_top
    c3.border = border

# ── Column Widths ──────────────────────────────────────────────────────────────
ws.column_dimensions['A'].width = 42
ws.column_dimensions['B'].width = 52
ws.column_dimensions['C'].width = 80

# ── Row Heights ────────────────────────────────────────────────────────────────
ws.row_dimensions[1].height = 28
for r in range(2, len(excel_rows) + 2):
    ws.row_dimensions[r].height = 30

# ── Freeze and Filter ──────────────────────────────────────────────────────────
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:C{len(excel_rows)+1}"

# ── Save ───────────────────────────────────────────────────────────────────────
out = "/home/user/Harrsh25/Substation_Workflow_Module_Guide.xlsx"
wb.save(out)
print(f"Saved: {out} | Data rows: {len(data)} | Excel rows: {len(excel_rows)}")
