import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Distribution Workflows"

# ── URL Aliases ────────────────────────────────────────────────────────────────
DIST_GEN    = "https://electrical-engineering-portal.com/distribution-substation"
ELEC4U_DIST = "https://electrical4u.com/electrical-power-distribution-system-in-power-systems/"
BID_EPC     = "https://constructionfront.com/epc-contract/"
FIDIC_URL   = "https://fidic.org/sites/default/files/The%20FIDIC%20Contracts%20Guide.pdf"
FIDIC_ACAD  = "https://www.academia.edu/8347063/FIDIC_Conditions_of_Contract_for_EPC_Turnkey_Projects_doc"
PMGR_URL    = "https://www.projectmanager.com/blog/managing-construction-projects"
PROC_URL    = "https://blog.projectmaterials.com/category/epc-projects/procurement-logistics/project-procurement-documents"
RKS_URL     = "https://rkstrainings.com/what-is-project-management-in-epc/"
SURV_DES    = "https://electrical-engineering-portal.com/power-substation-design-engineering"
KEEN_GUIDE  = "https://keentelengineering.com/a-guide-to-the-substation-design-process"
OHL_CONST   = "https://electrical-engineering-portal.com/power-substation-project-design-construction-erection-commissioning"
SUBSTR_FAULTS = "https://www.substationfaults.com/electrical-substation-construction/"
ELEC4U_OHL  = "https://electrical4u.com/overhead-line-conductors/"
ELEC4U_ELDIST = "https://electricaltechnology.org/2014/03/electrical-power-distribution-system.html"
DTR_FAT     = "https://electrical-engineering-portal.com/transformer-factory-acceptance-tests-final-inspections"
ELEC4U_DTR  = "https://electrical4u.com/distribution-transformer/"
KEEN_TEST   = "https://keentelengineering.com/substation-testing-lifecycle-fat-vs-sat"
CABLE_LAY   = "https://electrical-engineering-portal.com/power-cable-laying-installation-methods"
EEPOWER_PRE = "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-2-pre-commissioning-inspections/"
ELEC4U_CABLE = "https://electrical4u.com/underground-cable-types-advantages-disadvantages/"
SCADA_INSTR = "https://instrumentationtools.com/scada-for-substation-automation/"
ELEC4U_METER = "https://electrical4u.com/energy-meter-types-of-energy-meter/"
EEPOWER_SCADA = "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-4-scada-and-comms-systems"
DIST_AUTO   = "https://electrical-engineering-portal.com/distribution-automation-system"
EEPOWER_SCOPE = "https://eepower.com/technical-articles/substation-commissioning-and-testingpart-1-scope-and-workflow/"
EEPOWER_FIELD = "https://eepower.com/technical-articles/substation-commissioning-and-testing-part-3-field-testing/"
ELEC4LEARN  = "https://electrical4learning.blogspot.com/2024/11/method-statement-for-testing-and-commissioning-of-substation-equipments.html"
PWRSYNCHRO  = "https://powersynchro.com/substation-commissioning-checklist/"
DLP_MASTT   = "https://www.mastt.com/blogs/defects-liability-period"
DLP_RELGROW = "https://relgrow.com/resources/dlp-full-form-in-construction/"
DLP_CONLAW  = "https://constructionlawmadeeasy.com/construction-law/chapter-10/defects-liability-period/"
HANDOVER    = "https://projectmanagement123.com/project-completion-and-handover-procedure/"
FAT_CHECK   = "https://electrical-engineering-portal.com/checks-for-successful-substation-factory-acceptance-testing"
ELEC4U_EARTH = "https://electrical4u.com/electrical-earthing/"
KEEN_EARTH  = "https://keentelengineering.com/ieee-compliant-ehv-hv-mv-substation-design-services-by-keentel-engineering"
AT_C_LOSS   = "https://electrical-engineering-portal.com/distribution-substation"
ELEC4U_TX   = "https://electrical4u.com/transmission-line-in-power-system/"

# ── Data: (module, sub_module_step, url1, url2) ───────────────────────────────
data = [
    # MODULE 1 — Bid Process & Tender Management
    ("Bid Process & Tender Management",
     "Tender / NIT Identification – Monitor govt portals, receive NIT, document purchase",
     BID_EPC, PROC_URL),

    ("",
     "Bid / No-Bid Decision – Scope review, eligibility check, competition analysis, resource check, Go/No-Go",
     BID_EPC, RKS_URL),

    ("",
     "Pre-Bid Activities – Pre-bid meeting, site visit, queries submission, corrigendum review",
     BID_EPC, PMGR_URL),

    ("",
     "Scope & BOQ Analysis – BOQ review, technical spec study, quantity validation, ambiguous items",
     SURV_DES, KEEN_GUIDE),

    ("",
     "Cost Estimation – Material cost, civil cost, labour cost, equipment cost, overheads, taxes",
     RKS_URL, PROC_URL),

    ("",
     "Vendor & Subcontractor Quotations – RFQ to vendors, subcontractor quotes, rate comparison",
     PROC_URL, FAT_CHECK),

    ("",
     "Risk & Contingency Analysis – ROW risk, soil risk, monsoon risk, regulatory risk, consumer coordination risk",
     BID_EPC, PMGR_URL),

    ("",
     "Bid Price Finalization – Base cost consolidation, margin addition, management approval",
     RKS_URL, BID_EPC),

    ("",
     "Bid Preparation & Submission – Technical bid, commercial bid, EMD arrangement, submission, bid opening",
     FIDIC_URL, BID_EPC),

    ("",
     "Post-Bid Activities – Techno-commercial clarifications, price negotiation, LOA received",
     PROC_URL, FIDIC_ACAD),

    # MODULE 2 — Contract Finalization
    ("Contract Finalization",
     "LOA Review & Acceptance – LOA terms review, scope confirmation, acceptance letter",
     FIDIC_URL, RKS_URL),

    ("",
     "Contract Drafting & Negotiation – Draft review, penalty/LD clause, variation clause, DLP terms, dispute resolution",
     FIDIC_URL, FIDIC_ACAD),

    ("",
     "Contract Signing – Final execution, stamp duty, notarization",
     FIDIC_URL, BID_EPC),

    ("",
     "Bank Guarantee Submission – PBG submission (5-10%), Security Deposit, BG format approval",
     FIDIC_ACAD, DLP_CONLAW),

    ("",
     "Insurance & Compliance – CAR insurance, Workmen Compensation, Third Party Liability, ESI/PF docs",
     PMGR_URL, RKS_URL),

    ("",
     "Contract Effective & Project Kickoff – Contract effective date, project schedule (L1/L2), internal handover",
     RKS_URL, PMGR_URL),

    # MODULE 3 — Project Mobilization
    ("Project Mobilization",
     "Financial Mobilization – Mobilization advance claim, advance BG submission, fund allocation",
     FIDIC_ACAD, DLP_CONLAW),

    ("",
     "Team Mobilization – PM deployment, site engineers/surveyors, billing engineer, safety officer, labour gangs",
     RKS_URL, PMGR_URL),

    ("",
     "Site Office Setup – Site office establishment, material store/yard, IT setup",
     OHL_CONST, SUBSTR_FAULTS),

    ("",
     "Equipment & Vehicles Mobilization – Boring machines, JCB/excavator, crane/hydra, vehicles, safety equipment",
     OHL_CONST, SUBSTR_FAULTS),

    ("",
     "HSE Setup – HSE plan preparation, safety induction, mock drills, incident reporting system",
     ELEC4U_DIST, ELEC4LEARN),

    ("",
     "Regulatory Compliances – Labour license, municipal registration, forest/highway/railway permissions",
     PMGR_URL, RKS_URL),

    # MODULE 4 — Detailed Survey & Route Finalization
    ("Detailed Survey & Route Finalization",
     "HT Line Route Survey – Walkover survey, GPS/GIS mapping, obstacle identification, crossing identification, span measurement, ROW verification, pole schedule",
     SURV_DES, KEEN_GUIDE),

    ("",
     "LT Line Route Survey – Consumer list verification, LT route marking, service connection point identification",
     ELEC4U_DIST, ELEC4U_ELDIST),

    ("",
     "Soil Investigation – Soil type identification, foundation type selection, water table assessment",
     SURV_DES, KEEN_GUIDE),

    ("",
     "DTR Location Survey – Transformer site selection, civil pad feasibility, proximity to load center, access road availability",
     ELEC4U_DTR, DIST_GEN),

    ("",
     "Underground Cable Route Survey – Existing utility mapping, road crossing identification (HDD vs open cut), manhole/joint bay planning, route marker planning",
     CABLE_LAY, ELEC4U_CABLE),

    ("",
     "Survey Report & Client Approval – Survey report compilation, BOM revision, submission to client, IFC release",
     SURV_DES, KEEN_GUIDE),

    # MODULE 5 — Detailed Engineering & Drawing Preparation
    ("Detailed Engineering & Drawing Preparation",
     "HT Line Design – Conductor sag-tension calculations, pole schedule finalization, pole loading calculations, earthing design, DO fuse/AB switch locations, line drawing (plan & profile)",
     OHL_CONST, ELEC4U_OHL),

    ("",
     "LT Network Design – Load calculation, conductor size selection (voltage drop), LT pole schedule, service connection schedule",
     ELEC4U_DIST, ELEC4U_ELDIST),

    ("",
     "DTR Design – DTR capacity selection, HT connection scheme, LT panel design, earthing scheme, civil pad drawing",
     ELEC4U_DTR, DTR_FAT),

    ("",
     "Underground Cable Design – Cable sizing (current capacity/voltage drop), cable route drawing with chainage, joint bay/termination schedule, trench cross-section drawing",
     CABLE_LAY, ELEC4U_CABLE),

    ("",
     "Smart Meter / AMI Design – Meter type selection, communication technology selection (RF/GPRS/PLC), HES architecture, DCU location planning, MDM software configuration",
     ELEC4U_METER, SCADA_INSTR),

    ("",
     "Distribution Automation Design – Auto Recloser location, sectionalizer/FRTU selection, FLISR logic design, communication network design, SCADA integration architecture",
     DIST_AUTO, EEPOWER_SCADA),

    ("",
     "Drawing Submission & Approval – Submission to DISCOM/client, review comments incorporation, IFC drawings released",
     SURV_DES, KEEN_GUIDE),

    # MODULE 6 — Procurement & Material Management
    ("Procurement & Material Management",
     "Procurement Planning – MRP from IFC drawings, long lead item identification (DTRs/smart meters/cables), procurement schedule",
     PROC_URL, RKS_URL),

    ("",
     "Vendor Identification & RFQ – DISCOM-approved vendor list, RFQ issuance, technical bid evaluation, commercial comparison",
     PROC_URL, FAT_CHECK),

    ("",
     "Purchase Order Issuance – PO preparation with specs, delivery schedule, TPI inspection clause",
     PROC_URL, RKS_URL),

    ("",
     "Key Materials Procurement – Poles (PCC/steel/MS), conductors (ACSR/ABC), DTRs (25-250 kVA), insulators, hardware fittings, AB switches, dropout fuses, lightning arrestors, LT distribution boxes, XLPE cables, smart meters, RMU, auto reclosers, earthing materials",
     OHL_CONST, ELEC4U_DTR),

    ("",
     "Vendor Inspection (TPI/FAT) – Factory inspection for DTRs (IS 2026 tests), type test certificate verification for poles/conductors, meter type approval/NABL certificates",
     DTR_FAT, KEEN_TEST),

    ("",
     "Dispatch & Logistics – Dispatch clearance after inspection, route planning for oversized loads, transit insurance, staggered dispatch as per construction program",
     PROC_URL, RKS_URL),

    # MODULE 7 — Material Receipt & Store Management
    ("Material Receipt & Store Management",
     "Material Receipt at Site – Delivery challan verification against PO, physical quantity verification, visual inspection, GRN preparation",
     PROC_URL, SUBSTR_FAULTS),

    ("",
     "Incoming Material Inspection – Pole inspection (length/class/crack), conductor inspection (drum condition/length/dia), DTR inspection (nameplate/oil level/bushing/silica gel), hardware/insulator inspection, cable drum inspection, smart meter inspection",
     DTR_FAT, KEEN_TEST),

    ("",
     "Short Supply / Damage Claim (SDR) – Shortage/damage report to vendor, replacement claim processing, insurance claim for transit damage",
     PROC_URL, DLP_CONLAW),

    ("",
     "Store Management – Material segregation by type/location, pole yard stacking, conductor drum storage (upright/off ground), DTR storage (level ground/secured), FIFO management, daily material issue register",
     PROC_URL, RKS_URL),

    ("",
     "Issue to Field – Material indent from field engineer, issue as per approved drawings, material tracking (issue vs installation reconciliation)",
     PROC_URL, OHL_CONST),

    # MODULE 8 — Civil Works
    ("Civil Works",
     "Pole Foundation (Overhead Lines) – Pole location marking per pole schedule, pit excavation by soil type (normal/rocky/waterlogged), PCC foundation base, backfilling and compaction, foundation concrete curing",
     OHL_CONST, SUBSTR_FAULTS),

    ("",
     "Stay Foundation – Stay pit excavation (angle/section/dead end poles), stay block installation, stay rod and stay wire arrangement",
     OHL_CONST, ELEC4U_OHL),

    ("",
     "DTR Civil Pad (Ground-Mounted) – Excavation and levelling, RCC pad construction with transformer rail channels, plinth protection, fence/compound wall, oil soak pit construction",
     ELEC4U_DTR, SUBSTR_FAULTS),

    ("",
     "Earth Pit Construction – Earth pit excavation (2-3m depth), GI pipe/plate electrode installation, salt and charcoal filling, earthing wire connection, earth resistance measurement",
     ELEC4U_EARTH, KEEN_EARTH),

    ("",
     "Trench Excavation (UG Projects) – Trench marking per cable route drawing, excavation by location (JCB/HDD/manual), trench depth verification, sand bedding (150mm), route marker brick laying, backfilling in layers with compaction",
     CABLE_LAY, ELEC4U_CABLE),

    ("",
     "Joint Bay / Manhole Construction – Joint bay excavation, RCC structure construction, cable entry/exit duct sealing, cover slab with lifting provision",
     CABLE_LAY, ELEC4U_CABLE),

    ("",
     "Control Room / Panel Room (DA/SCADA) – Prefabricated/brick structure construction, cable entry conduits, earthing and lightning protection",
     DIST_AUTO, SCADA_INSTR),

    # MODULE 9 — Pole Erection & Overhead Line Construction
    ("Pole Erection & Overhead Line Construction",
     "Pole Erection – Pole transport to location, pole erection using crane/hydra (vertical alignment/rake check), setting in foundation pit and backfilling, compaction and stability check, cross arm and insulator pin fixing",
     OHL_CONST, SUBSTR_FAULTS),

    ("",
     "Cross Arm & Hardware Fitting – Cross arm fixing (11kV configuration), GOI/AB switch mounting, dropout fuse mounting, lightning arrestor mounting, stay wire fitting",
     OHL_CONST, ELEC4U_OHL),

    ("",
     "Insulator Fitting – Pin insulator for tangent poles, strain/disc insulator assembly for dead-end/angle poles, visual inspection before fitting",
     ELEC4U_OHL, ELEC4U_ELDIST),

    ("",
     "HT Conductor Stringing (11/22/33kV) – Conductor drum placement, pilot wire/pulling rope laying, conductor pulling (manual/mechanical winch), sagging (sag board/dynamometer method), conductor binding/clamping at insulators, jumper connections",
     ELEC4U_OHL, OHL_CONST),

    ("",
     "LT Conductor / ABC Cable Stringing (415V) – LT cross arm/bracket fixing, ABC cable stringing with messenger wire tensioning, strain/suspension clamp fitting, service connection tee-off joints",
     ELEC4U_DIST, ELEC4U_ELDIST),

    ("",
     "Conductor Joints – Mid-span compression joints, parallel groove clamps for T-off connections, ferrule joints for small conductors",
     ELEC4U_OHL, OHL_CONST),

    # MODULE 10 — DTR Installation
    ("DTR Installation",
     "DTR Transportation to Site – Proper blocking on truck, tilt indicator check, site accessibility check",
     ELEC4U_DTR, PROC_URL),

    ("",
     "DTR Erection (Pole-Mounted) – Transformer brackets/cradle fitting on poles, lifting with hydra/crane, securing on brackets",
     ELEC4U_DTR, OHL_CONST),

    ("",
     "DTR Erection (Ground-Mounted) – Placement on RCC pad with rollers, alignment with rail channels, securing with anchor bolts",
     ELEC4U_DTR, SUBSTR_FAULTS),

    ("",
     "HT Connections – HT AB switch/dropout fuse wiring from feeder line to HV bushing, lightning arrestor connection, HT cable/jumper termination at HV bushing",
     ELEC4U_DTR, ELEC4U_OHL),

    ("",
     "LT Connections – LT bus bar/link connection from LV bushing, LT distribution box/metering panel installation, LT outgoing feeder connections",
     ELEC4U_DTR, ELEC4U_DIST),

    ("",
     "DTR Earthing – Body/tank earthing (GI flat to earth electrode), neutral earthing (LV neutral to earth), LA earthing (separate earth pit), earth resistance measurement",
     ELEC4U_EARTH, KEEN_EARTH),

    ("",
     "Oil Treatment & Filling – Oil BDV testing (minimum 30 kV/2.5mm), oil filtration if BDV low, oil level check",
     DTR_FAT, ELEC4U_DTR),

    ("",
     "DTR Pre-Commissioning – IR test (HV to LV/HV to earth/LV to earth), Turns Ratio Test (TTR), vector group verification, no-load losses measurement, continuity check",
     DTR_FAT, KEEN_TEST),

    # MODULE 11 — Underground Cable Laying & Termination
    ("Underground Cable Laying & Termination",
     "Cable Pulling – Cable drum positioning, cable laying in trench (direct buried/in duct), cable bend radius compliance, route marker tiles laying, warning tape laying, cable length reconciliation",
     CABLE_LAY, ELEC4U_CABLE),

    ("",
     "Cable Jointing – Joint location identification, cold shrink/heat shrink joint kit installation (cable stripping/conductor connection/stress cone/heat shrink tube), joint bay sealing, joint record (GPS/depth/type)",
     CABLE_LAY, ELEC4U_CABLE),

    ("",
     "Cable Termination – Indoor/outdoor termination kit installation, cable glanding at RMU/panel, armour earthing at termination, stress cone installation, cable tag and phase identification",
     CABLE_LAY, EEPOWER_PRE),

    ("",
     "High Voltage Testing of Cables – IR test before and after laying, VLF hipot test (0.1 Hz, 17kV for 60 min for 11kV cable), partial discharge test for critical circuits, sheath continuity and insulation test",
     EEPOWER_PRE, EEPOWER_FIELD),

    ("",
     "RMU Installation – RMU positioning and levelling, cable termination at incoming/outgoing feeders, earthing of RMU enclosure, SF6 gas pressure check, functional test (open/close operations)",
     ELEC4U_CABLE, EEPOWER_PRE),

    # MODULE 12 — Smart Meter / AMI Installation
    ("Smart Meter / AMI Installation",
     "Meter Hardware Installation – Existing meter removal, smart meter mounting on consumer meter board, CT installation (3-phase/HT consumers), tamper-evident sealing, consumer detail tagging",
     ELEC4U_METER, SCADA_INSTR),

    ("",
     "Communication Module Setup – SIM card/RF module insertion and configuration, DCU installation at feeder pillar/DTR, communication test (meter to DCU), DCU to HES link verification",
     SCADA_INSTR, EEPOWER_SCADA),

    ("",
     "HES / MDM Configuration – Meter registration in HES with consumer data, tariff loading in meter, time sync (meter clock synchronization), data pull test, alert configuration (tamper/power fail/reverse current)",
     SCADA_INSTR, EEPOWER_SCADA),

    ("",
     "Consumer Onboarding – Consumer information update in DISCOM system, opening reading recording, consumer awareness/demo",
     ELEC4U_METER, DIST_GEN),

    # MODULE 13 — Distribution Automation (DA) & SCADA Integration
    ("Distribution Automation (DA) & SCADA Integration",
     "Field Device Installation – Auto Recloser installation (pole mounting/HV cable connection/control cable to RTU), sectionalizer installation, FRTU installation in control box, CT/PT metering unit installation, solar panel/battery backup for FRTU",
     DIST_AUTO, SCADA_INSTR),

    ("",
     "Communication Network Setup – GPRS modem configuration in FRTU, optical fiber cable laying if applicable, RF mesh network setup, communication test (FRTU to SCADA connectivity)",
     SCADA_INSTR, EEPOWER_SCADA),

    ("",
     "SCADA / DMS Configuration – Network topology loading in SCADA (feeder diagram/SLD/equipment tagging), RTU/FRTU protocol configuration (DNP3/IEC 60870-5-101/104), control function testing (remote open/close), FLISR logic testing, alarm and event configuration",
     DIST_AUTO, EEPOWER_SCADA),

    ("",
     "Integration Testing – End-to-end test (field trip → SCADA alarm → auto isolation → restoration), communication redundancy test, SCADA data accuracy verification",
     SCADA_INSTR, EEPOWER_SCADA),

    # MODULE 14 — Electrical Testing & Pre-Commissioning
    ("Electrical Testing & Pre-Commissioning",
     "HT Line Testing – IR test (phase to phase/phase to earth), continuity test, phase sequence verification, visual inspection (hardware tightness/sag/clearances), minimum ground clearance verification (6.1m over roads, railway norms)",
     EEPOWER_PRE, ELEC4LEARN),

    ("",
     "LT Line Testing – IR test of LT network, phase identification at DTR LV terminals, LT feeder polarity check",
     EEPOWER_PRE, ELEC4U_DIST),

    ("",
     "DTR Pre-Commissioning Tests – IR test (all combinations), Turns Ratio Test, no load test (no-load current), oil BDV test (minimum 30 kV), earth resistance measurement",
     DTR_FAT, KEEN_TEST),

    ("",
     "AB Switch / DO Fuse Testing – Mechanical operation test (open/close 5 times), contact resistance measurement, insulation resistance test",
     EEPOWER_PRE, EEPOWER_FIELD),

    ("",
     "Underground Cable Testing – IR test before and after laying, VLF hipot test, sheath continuity test, phase identification at both ends",
     EEPOWER_FIELD, ELEC4U_CABLE),

    ("",
     "Earthing System Verification – Earth resistance measurement at each DTR (per IE Rules), neutral earthing continuity check, body/equipment earthing continuity check",
     ELEC4U_EARTH, KEEN_EARTH),

    ("",
     "Punch List Preparation – Pre-commissioning inspection punch list, closure of punch list items, pre-commissioning completion certificate",
     PWRSYNCHRO, EEPOWER_SCOPE),

    # MODULE 15 — Commissioning & Energization
    ("Commissioning & Energization",
     "External Clearances – Commissioning application to DISCOM/regulatory authority, Electrical Inspector (EI) inspection request, EI inspection and certificate issuance, DISCOM PTCC (Permission to Charge Certificate)",
     EEPOWER_SCOPE, ELEC4LEARN),

    ("",
     "Pre-Energization Safety Check – Safety clearance (all workers off line), earthing clamps removed from all phases, all hardware secured, safety barriers/caution tape, switching scheme briefing",
     EEPOWER_SCOPE, ELEC4LEARN),

    ("",
     "HT Line Energization – Step-by-step energization from source end, first charge observation (flashovers/tripping), voltage measurement at receiving end, phase sequence verification at all DTR HV terminations",
     EEPOWER_SCOPE, ELEC4LEARN),

    ("",
     "DTR Energization – HT fuse/AB switch closing (DTR charging), no-load voltage measurement on LT side (415V ±5%), phase sequence check LT side, neutral voltage measurement",
     ELEC4U_DTR, EEPOWER_SCOPE),

    ("",
     "LT Network Energization – LT feeder links closing one by one, voltage measurement at feeder end (voltage drop check), load connection feeder by feeder, load current measurement at DTR LT terminals",
     ELEC4U_DIST, EEPOWER_SCOPE),

    ("",
     "Smart Meter Commissioning – Meter power-on verification, real-time energy reading display check, communication link active status in HES, tamper seal verification",
     ELEC4U_METER, SCADA_INSTR),

    ("",
     "DA / SCADA Live Commissioning – Auto Recloser live operation test, SCADA real-time data verification, remote switching test from SCADA",
     DIST_AUTO, EEPOWER_SCADA),

    # MODULE 16 — Trial Run & Performance Monitoring
    ("Trial Run & Performance Monitoring",
     "Trial Operation Period – System running under normal load (30 days), daily monitoring (voltage/load current/power factor), outage recording (fault type/location/restoration time), smart meter data collection",
     EEPOWER_SCOPE, ELEC4LEARN),

    ("",
     "AT&C Loss Measurement – DTR metering (energy input at DTR), consumer metering (energy billed), AT&C loss calculation, loss report submission to client",
     AT_C_LOSS, ELEC4U_TX),

    ("",
     "Snag Rectification – Punch list from trial run, voltage problem rectification (tap change/conductor upgrade), loose connection tightening, earthing improvement",
     PWRSYNCHRO, EEPOWER_PRE),

    ("",
     "Performance Stabilization Certificate – Trial run completion report, performance stabilization confirmed by client",
     EEPOWER_SCOPE, PMGR_URL),

    # MODULE 17 — Project Handover & Closure
    ("Project Handover & Closure",
     "As-Built Documentation – As-built drawings (pole schedule/route map/GIS data), test certificates (IR/earth resistance/VLF/oil BDV), material reconciliation (issued vs installed vs balance), consumer list with meter numbers, O&M manuals",
     HANDOVER, PMGR_URL),

    ("",
     "Client Walk-Through & Acceptance – Joint inspection with client/DISCOM team, punch list review and closure confirmation, commissioning report acceptance",
     HANDOVER, PMGR_URL),

    ("",
     "Provisional Acceptance Certificate (PAC) – PAC issuance by client, retention money partial release",
     DLP_RELGROW, DLP_MASTT),

    ("",
     "Financial Closure – Final bill preparation with all variation orders, measurement book reconciliation, final bill submission and approval, payment receipt",
     FIDIC_ACAD, FIDIC_URL),

    ("",
     "Regulatory Handover – Handover to DISCOM O&M team, line details in DISCOM GIS/asset management, consumer data transfer to DISCOM billing system",
     HANDOVER, DIST_GEN),

    # MODULE 18 — Defect Liability Period (DLP)
    ("Defect Liability Period (DLP)",
     "DLP Period Management – DLP period tracking (12-24 months from PAC), periodic site visits, defect register maintenance",
     DLP_MASTT, DLP_RELGROW),

    ("",
     "Defect Attendance – Defect notification from client, defect investigation (root cause analysis), repair/replacement execution (conductor snap/insulator failure/DTR failure/smart meter failure), defect closure confirmation",
     DLP_MASTT, DLP_CONLAW),

    ("",
     "Warranty Claims Management – Back-to-back warranty claim to vendor, DTR warranty replacement (12-24 months), smart meter warranty replacement",
     DLP_MASTT, FIDIC_ACAD),

    # MODULE 19 — Final Acceptance Certificate (FAC) & Contract Closeout
    ("Final Acceptance Certificate (FAC) & Contract Closeout",
     "FAC Application – All DLP defects closed confirmation, FAC application submission",
     DLP_RELGROW, DLP_MASTT),

    ("",
     "FAC Issuance – Final joint inspection, FAC issued by client",
     DLP_RELGROW, DLP_CONLAW),

    ("",
     "Retention Release – PBG/BG return by client, retention money release, final payment receipt",
     DLP_CONLAW, FIDIC_ACAD),

    ("",
     "Contract Closure – Contract closure letter exchange, internal project closure (lessons learned document), project archive (all documents stored)",
     HANDOVER, PMGR_URL),
]

# ── Build Workbook ─────────────────────────────────────────────────────────────
thin = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
border_no_right = Border(left=thin, right=Side(style=None), top=thin, bottom=thin)
border_no_left  = Border(left=Side(style=None), right=thin, top=thin, bottom=thin)

hdr_font  = Font(bold=True, size=11)
hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Columns A, B headers
for col, h in enumerate(["Module", "Sub-Module / Step"], 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = hdr_font
    cell.alignment = hdr_align
    cell.border = border

# Merged header C1:D1 — "Source URLs"
ws.merge_cells("C1:D1")
hdr_url = ws.cell(row=1, column=3, value="Source URLs")
hdr_url.font = hdr_font
hdr_url.alignment = hdr_align
hdr_url.border = border

wrap_top = Alignment(vertical='top', wrap_text=True)

for row_num, (mod, sub, url1, url2) in enumerate(data, 2):
    # Col A — Module
    c1 = ws.cell(row=row_num, column=1, value=mod)
    c1.alignment = wrap_top
    c1.border = border
    if mod:
        c1.font = Font(bold=True)

    # Col B — Sub-Module
    c2 = ws.cell(row=row_num, column=2, value=sub)
    c2.alignment = wrap_top
    c2.border = border

    # Col C — Source URL 1 (clickable, no right border)
    c3 = ws.cell(row=row_num, column=3, value=url1)
    c3.hyperlink = url1
    c3.font = Font(color="0563C1", underline="single")
    c3.alignment = wrap_top
    c3.border = border_no_right

    # Col D — Source URL 2 (clickable, no left border)
    c4 = ws.cell(row=row_num, column=4, value=url2)
    c4.hyperlink = url2
    c4.font = Font(color="0563C1", underline="single")
    c4.alignment = wrap_top
    c4.border = border_no_left

# ── Column Widths ──────────────────────────────────────────────────────────────
ws.column_dimensions['A'].width = 42
ws.column_dimensions['B'].width = 52
ws.column_dimensions['C'].width = 62
ws.column_dimensions['D'].width = 62

# ── Row Heights ────────────────────────────────────────────────────────────────
ws.row_dimensions[1].height = 28
for r in range(2, len(data) + 2):
    ws.row_dimensions[r].height = 40

# ── Freeze and Filter ──────────────────────────────────────────────────────────
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:D{len(data)+1}"

# ── Save ───────────────────────────────────────────────────────────────────────
out = "/home/user/Harrsh25/Distribution_Workflow_Module_Guide.xlsx"
wb.save(out)
print(f"Saved: {out} | Rows: {len(data)}")
