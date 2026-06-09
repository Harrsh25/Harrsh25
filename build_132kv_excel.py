import csv
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

CSV_PATH = "/root/.claude/uploads/88ea236d-50b0-5ace-9b45-3e67282b94dc/7767431f-132_KV.csv"
OUT_PATH = "/home/user/Harrsh25/132_KV_Transmission_WBS.xlsx"

HEADERS = [
    "Level1_Name","Level1_Desc",
    "Level2_Name","Level2_Desc",
    "Level3_Name","Level3_Desc",
    "Level4_Name","Level4_Desc",
    "Level5_Name","Level5_Desc",
    "Assign_To","Timeline","Weight","Tower_Tagging","Leaf_Node_Progress_Type"
]

T = Side(style="thin", color="000000")
M = Side(style="medium", color="000000")

def bdr(thick=False):
    s = M if thick else T
    return Border(left=s, right=s, top=s, bottom=s)

def fnt(bold=False, sz=10):
    return Font(bold=bold, size=sz, name="Calibri", color="000000")

def aln(h="left", wrap=True):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap)

def no_fill():
    return PatternFill(fill_type=None)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "132 KV Transmission WBS"

# Write header row
for col_idx, hdr in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col_idx, value=hdr)
    cell.font = fnt(bold=True, sz=10)
    cell.alignment = aln(h="center", wrap=True)
    cell.border = bdr(thick=True)
    cell.fill = no_fill()

# Read CSV and write data rows
with open(CSV_PATH, newline='', encoding='utf-8') as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, 2):
        # Pad row to 15 columns
        while len(row) < 15:
            row.append("")
        
        # Determine level depth for bold/size styling
        # Level1 = col A (index 0), Level3 = col E (index 4), Level5 = col I (index 8)
        level = 0
        if row[0].strip():
            level = 1
        elif row[2].strip():
            level = 2
        elif row[4].strip():
            level = 3
        elif row[6].strip():
            level = 4
        elif row[8].strip():
            level = 5
        
        bold = level <= 2
        sz = 11 if level == 1 else (10 if level == 2 else 9)
        
        for col_idx, val in enumerate(row[:15], 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val.strip() if val else "")
            cell.font = fnt(bold=bold, sz=sz)
            cell.alignment = aln()
            cell.border = bdr()
            cell.fill = no_fill()

# Set column widths
col_widths = [25, 30, 25, 30, 25, 30, 25, 30, 25, 30, 18, 25, 8, 14, 22]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# Freeze header row
ws.freeze_panes = "A2"

wb.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
