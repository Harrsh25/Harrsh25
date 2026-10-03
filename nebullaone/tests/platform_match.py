# Feature-by-feature audit: how NebullaOne works vs ERPNext, Odoo, Zoho Books, Oracle Fusion, SAP S/4HANA (+ Ariba)
# Usage: python3 tests/platform_match.py [out.xlsx]
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'docs', 'Platform_Match_Audit.xlsx')
M, F, A, P, R = 'Matches', 'Fixed now — matches', 'Ahead', 'Prototype limit', 'Removed at your request'
H = ['Module', 'Feature', 'How NebullaOne works', 'How the platforms work', 'Status', 'What was changed / why']
rows = [
 # Vendor master
 ('Vendor master', 'Quick create, rest later', 'Short Register vendor form (~10 fields); "Fill all details now" for the full form', 'ERPNext / Odoo / Zoho / SAP: short create, record filled later', M, ''),
 ('Vendor master', 'Vendor completes own registration', 'Invite (vendor fills everything) or Save & send to vendor (vendor completes your draft, same record)', 'Oracle supplier registration, Ariba SLP: supplier fills, buyer approves', M, ''),
 ('Vendor master', 'What the vendor supplies', 'Goods / Services / Labour in any combination; Labour or Services makes a contractor', 'Oracle / Ariba: many product & service categories per supplier', M, 'Built in the last rounds'),
 ('Vendor master', 'Trades / categories', 'Several per vendor, grouped, searchable dropdown; drive RFQ invitations', 'Oracle products & services, SAP Ariba UNSPSC commodities (multi, drive sourcing)', M, ''),
 ('Vendor master', 'GSTIN / PAN with duplicate check', 'GSTIN fills PAN and state; duplicate GSTIN / PAN flagged', 'ERPNext India, Zoho India: GSTIN validation, duplicates warned', M, ''),
 ('Vendor master', 'Registration tier (Prospective / Spend Authorized)', 'Prospective can quote, cannot get PO / contract; upgrade needs Finance approval', 'Oracle: Prospective vs Spend Authorized supplier', M, ''),
 ('Vendor master', 'Supplier tier / segmentation', 'Strategic / Preferred / Approved / Transactional; Strategic gets yearly questionnaire and extra compliance', 'SAP Ariba segmentation (manual, by category manager), Oracle supplier management', M, 'Segmentation is a manual decision in the platforms too'),
 ('Vendor master', 'Vendor group', 'Tree (Parent › Child); filters include children; group default payment terms fill the vendor', 'ERPNext Supplier Group tree with default payment terms', F, 'Added payment terms per group (Procurement Settings → Masters); picking the group fills the terms'),
 ('Vendor master', 'Internal / group company', 'Marked as group company; PO notes inter-company, no RFQ needed', 'ERPNext internal supplier, SAP trading partner', M, 'Inter-company accounting entries belong to the finance system'),
 ('Vendor master', 'Default buyer', 'Fills the Buyer on new POs for the vendor', 'Odoo purchase representative on the vendor', F, 'Was stored only; now fills the PO buyer'),
 ('Vendor master', 'Payment method', 'Preselects the payment mode when paying the vendor', 'Oracle / SAP / Odoo payment method on the supplier', F, 'Was stored only; lists unified with payment modes'),
 ('Vendor master', 'Receipt reminder (days)', 'N days before delivery the vendor is reminded automatically; shown on the PO and in the audit log', 'Odoo receipt reminder e-mail', F, 'Was stored only; now sent automatically'),
 ('Vendor master', 'Also a transporter', 'Transporter vendors are offered on goods receipts', 'ERPNext "Is Transporter" on the supplier', F, 'Was stored only; now a pick-list on the GRN'),
 ('Vendor master', 'Bill delivery', '—', 'Not a field in any of the five', F, 'Removed (did nothing, no platform equivalent)'),
 ('Vendor master', 'Credit limit', 'Warning on new POs when outstanding + PO exceeds it', 'Oracle / SAP credit checks (mainly customer side)', M, ''),
 ('Vendor master', 'Price list', 'Fills PO rates from the vendor price list', 'ERPNext default price list, Odoo vendor pricelist', M, ''),
 ('Vendor master', 'Hold / block with scope and release date', 'Hold for All / POs / payments with reason and release date; auto-release on the date', 'ERPNext hold (All / Invoices / Payments) with release date', M, ''),
 ('Vendor master', 'Bank account verification', 'Re-enter account, IFSC check, verified status, payments-enabled flag', 'Oracle / SAP bank validation, ERPNext bank account', M, ''),
 ('Vendor master', 'Tax extras (MSME, CIN, D-U-N-S, entity type)', '—', 'Zoho / ERPNext India (MSME), Oracle / SAP (registry IDs)', R, 'Removed from the form at your request'),
 ('Vendor master', 'Tags, logo, custom fields, notes', '—', 'Odoo tags, ERPNext notes', R, 'Removed at your request; comments on the record cover notes'),
 # Approvals & compliance
 ('Approvals', 'Multi-stage vendor approval', 'Configurable stages (Procurement → Legal → Finance), request changes, resubmit', 'Oracle approval rules, Ariba SLP approval flow, ERPNext workflow', M, ''),
 ('Approvals', 'Roles & who-can-approve checks', 'Anyone can act', 'All five enforce roles and segregation of duties', R, 'Removed at your request — needed before real users'),
 ('Compliance', 'Required documents, expiry, verification', 'Per vendor type; versions; expiry reminders; payment gate', 'Oracle / Ariba compliance; ERPNext none', A, ''),
 ('Compliance', 'Insurance cover minimums', 'Policy entry, cover checks, gate', 'Not standard in the five', A, ''),
 ('Qualification', 'Questionnaires, scoring, requalification', 'Rule sets + question library, score, limits, expiry, requalification queue', 'Oracle Supplier Qualification, Ariba SLP', M, ''),
 # Sourcing
 ('Requisition', 'Requisition with approval', 'Purpose, items, required by, approve / reject, % ordered / received', 'ERPNext Material Request, Oracle / SAP requisitions', M, ''),
 ('Requisition', 'Requisition → PO directly', 'Create PO from an approved requisition (lines, project, need-by carried)', 'ERPNext MR → PO, Oracle / SAP requisition to order', F, 'Was only via RFQ; now direct too'),
 ('Requisition', 'Requisition → RFQ', 'Create RFQ from the requisition', 'All five', M, ''),
 ('RFQ', 'Send RFQ, vendor quotes online', 'E-mail with secure link, one-time code, line quotes, no-bid, CSV, questions', 'Oracle / Ariba sourcing, ERPNext supplier portal', M, ''),
 ('RFQ', 'Cancel an RFQ', 'Cancel with reason while nothing is awarded; quotes stay on record', 'All five can cancel', F, 'Status existed but no action — added'),
 ('RFQ', 'Compare & award by line', 'Weighted ranking, award per line, one PO per vendor, award as contract', 'Oracle / Ariba award by line; ERPNext comparison report', M, ''),
 ('RFQ', 'Reverse auctions', '—', 'Oracle Sourcing, SAP Ariba', P, 'Not built; add only if you run auctions'),
 ('Blanket order', 'Agreement, call-offs, consumption', 'Agreed lines & rates, call-off POs, allowance, deadline, close', 'Odoo blanket orders, Oracle BPA, SAP outline agreements', M, ''),
 # PO
 ('Purchase order', 'Approval threshold', 'Below threshold issued directly, above goes for approval', 'All five', M, ''),
 ('Purchase order', 'Reject a draft PO', 'Reject with reason', 'All five', F, 'Function existed, no button — added'),
 ('Purchase order', 'Cancel PO', 'Allowed while nothing received or billed; stays on the list as Cancelled', 'ERPNext cancel, Odoo cancel, Oracle / SAP cancel', F, 'Added; cancelled POs no longer hidden'),
 ('Purchase order', 'Close / short-close', 'Partly received PO closed; rest not expected; received part still billable', 'ERPNext Close, Oracle Close, SAP delivery completed', F, 'Added'),
 ('Purchase order', 'Amend with revision history', 'Edit creates a revision; locked confirmed orders', 'ERPNext amend, Oracle change order, SAP versioning', M, ''),
 ('Purchase order', 'Line tax, HSN, need-by date', 'Per line', 'All five', M, 'Added in the field round'),
 ('Goods receipt', 'Receipt with QC, accepted / rejected, returns', 'Inspection template, rejected store, auto return + debit note', 'All five', M, ''),
 ('Goods receipt', 'Reverse a receipt', 'Allowed if not billed and nothing returned; kept as reversed', 'ERPNext cancel PR, Odoo return, SAP movement 102', F, 'Added'),
 # Bills & payments
 ('Bills', '3-way match with tolerances, holds', 'PO × GRN × bill with qty / amount tolerance; holds with release date', 'Oracle / SAP / ERPNext', M, ''),
 ('Bills', 'Cancel a bill', 'Allowed when unpaid; stops being payable, quantities billable again', 'ERPNext cancel, Zoho void, Oracle cancel, SAP reverse', F, 'Added'),
 ('Bills', 'Credit / debit notes', 'Linked to the bill', 'All five', M, ''),
 ('Bills', 'TDS per bill', 'Category per bill, threshold, multi-type vendors default by bill', 'ERPNext tax withholding, Zoho TDS per bill', M, ''),
 ('Payments', 'Batch payment, advances, write-off', 'Select payable, advance adjustment, write-off', 'All five', M, ''),
 ('Payments', 'Reverse a payment', 'Reverse with reason (bounced cheque…); balance returns', 'ERPNext cancel payment, Zoho delete, Odoo / SAP reverse', F, 'Added'),
 ('Payments', 'Real bank file / e-invoice', 'Simulated', 'Oracle / SAP / Zoho / ERPNext India', P, 'Needs bank / GST integration'),
 # Contract & labour
 ('Contract', 'Works contract with BOQ, approval, BG, retention, LD, DLP', 'Full lifecycle incl. guarantees register', 'Oracle complex work, SAP SES & retention cover parts', A, ''),
 ('Contract', 'Change orders, EOT, terminate, close', 'Value & quantity lines, time extension, termination → encash → final account', 'Oracle contract amendments, SAP change', A, ''),
 ('Work order', 'Issue, contractor accepts, suspend, short-close', 'Portal acceptance, mobilisation gate', 'SAP service PO / Oracle complex work', A, ''),
 ('Work order', 'Cancel a work order', 'Draft or issued with nothing measured', 'SAP / Oracle cancel the service order', F, 'Status existed, no action — added'),
 ('Measurement', 'Measurement book + JMS', 'Dimensioned entries, contractor co-sign, inspection, attendance roll-up', 'SAP service entry sheet, Oracle work confirmation', A, ''),
 ('Measurement', 'Withdraw a wrong entry', 'Unbilled entries can be withdrawn (kept in the log)', 'SAP revoke SES, Oracle withdraw work confirmation', F, 'Added'),
 ('RA bill', 'Progress bill with retention, advance, TDS, cess', 'Certification flow, deductions, reject', 'Not standard in ERPNext / Odoo / Zoho', A, ''),
 ('Closure', 'Close-out, DLP, final settlement, release, termination', 'Punch list → handover → DLP → final account → release certificate', 'Not standard in the five', A, ''),
 # Cross-cutting
 ('Scorecard', 'Supplier scorecard with standings', 'Formula criteria, weights, standings that warn / stop RFQ & PO, CAPs', 'ERPNext supplier scorecard (same idea)', M, ''),
 ('Portal', 'Supplier / contractor portal', 'RFQs, POs, WOs, RA claims, attendance, punch list, bills, documents', 'Oracle / Ariba / Zoho / ERPNext portals (no contractor screens)', A, ''),
 ('Audit', 'Audit trail', 'Every action, user, time, reason', 'All five', M, ''),
 ('Records', 'Comments with @mention', 'On every record, logged', 'Odoo chatter, ERPNext comments', M, ''),
 ('Lists', 'Search, filters, columns, board, calendar, export, paging', 'Quick filter icons + Filters panel', 'All five', M, 'Saved views & group-by removed at your request'),
 ('Accounting', 'GL postings, stock valuation', 'Not kept', 'All five post to the ledger', P, 'Belongs to your ERP / Tally'),
 ('E-mail', 'Real e-mails', 'Simulated (logged)', 'All five send', P, 'Needs a mail server'),
]
wb = Workbook(); ws = wb.active; ws.title = 'Audit'; ws.append(H)
fill = {M: 'E8F5E9', F: 'E3F2FD', A: 'EDE7F6', P: 'FFF3E0', R: 'FCE4EC'}
for r in rows:
    ws.append(list(r)); ws.cell(ws.max_row, 5).fill = PatternFill('solid', fgColor=fill[r[4]])
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79')
for i, w in enumerate([16, 34, 60, 52, 22, 52], 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
s = wb.create_sheet('Summary'); s.append(['Status', 'Features', 'Meaning'])
mean = {M: 'Works the same way as the platforms', F: 'Was partly / not working — fixed in this round', A: 'NebullaOne goes further (construction-specific)', P: 'Needs a real backend / integration, not possible in a single HTML file', R: 'You asked to remove it'}
for k in (M, F, A, P, R): s.append([k, sum(1 for r in rows if r[4] == k), mean[k]])
s.append(['Total', len(rows), ''])
for c in s[1]: c.font = Font(bold=True)
for i, w in enumerate([24, 10, 80], 1): s.column_dimensions[get_column_letter(i)].width = w
os.makedirs(os.path.dirname(OUT), exist_ok=True); wb.save(OUT)
print(OUT, len(rows), {k: sum(1 for r in rows if r[4] == k) for k in (M, F, A, P, R)})
