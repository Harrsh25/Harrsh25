# Builds docs/Workflow_UAT_and_Platform_Comparison.xlsx:
#   Summary · Workflow UAT (final.js, 33 lifecycle steps) · Field mapping (mapping.js) ·
#   Platform comparison (NebullaOne vs ERPNext, Odoo, Zoho, Oracle Fusion, SAP S/4HANA + Ariba) ·
#   Gaps & recommendations · Sources
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'docs', 'Workflow_UAT_and_Platform_Comparison.xlsx')
load = lambda n: json.load(open(os.path.join(HERE, 'out', f'res-{n}.json')))

HEAD = PatternFill('solid', fgColor='1F3A5F'); HF = Font(bold=True, color='FFFFFF')
thin = Side(style='thin', color='D7DCE3'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {'PASS': 'E3F4E8', 'FAIL': 'FBE3E3', 'Native': 'E3F4E8', 'Partial': 'FFF2D6', 'Add-on': 'EEF1F5', 'Not native': 'FBE3E3', 'Fixed': 'E3F4E8',
        'Ahead': 'D9EAFB', 'On par': 'E3F4E8', 'Behind': 'FBE3E3', 'Different': 'FFF2D6'}

def sheet(wb, title, cols, rows, widths, colour_cols=(), height=None):
    ws = wb.create_sheet(title); ws.append(cols)
    for c in ws[1]: c.fill = HEAD; c.font = HF; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True); c.border = BOX
            if c.column - 1 in colour_cols:
                key = next((k for k in FILL if str(c.value or '').startswith(k)), None)
                if key: c.fill = PatternFill('solid', fgColor=FILL[key])
        if height: ws.row_dimensions[row[0].row].height = height
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'B2' if len(cols) > 4 else 'A2'; ws.auto_filter.ref = ws.dimensions
    return ws

# ---------------------------------------------------------------- platform comparison
# Status words: Native = standard feature · Partial = covered in part / by configuration · Add-on = paid add-on, separate product or marketplace app · Not native
# Verdict = NebullaOne compared with the five platforms.
CMP = [
 # area, step, NebullaOne screen, NebullaOne has, ERPNext, Odoo, Zoho, Oracle, SAP, verdict, note
 ('A Onboarding', 'Vendor invitation & self-registration', 'Vendor Registry → Invite vendor; public /vendor-register',
  'Invite by e-mail, self-registration form with documents, bank, contacts; draft → submit',
  'Partial — supplier created by staff; portal user via RFQ / contact', 'Partial — vendor created by staff; portal access to contacts', 'Partial — vendor created by staff; vendor portal invite',
  'Native — Supplier registration (internal & external) with approval', 'Native — SAP Ariba SLP supplier requests & registration', 'On par', 'Matches Oracle / Ariba-style registration; ahead of ERPNext / Odoo / Zoho'),
 ('A Onboarding', 'Multi-stage vendor approval', 'Vendor Approvals; Procurement Settings → Gates & approvals',
  'Configurable stages (Procurement → Legal → Finance), scope per contractor / non-contractor, request changes, resubmit',
  'Partial — generic Workflow builder', 'Partial — approval rules / Studio', 'Partial — custom approvals', 'Native — approval rules on registration', 'Native — Ariba approval flows', 'On par', ''),
 ('A Onboarding', 'Qualification questionnaire, rule sets, requalification', 'Vendor record → Qualification tab; Requalification page',
  'Rule-set questions + question library, score, limits (single / aggregate), expiry, risk; requalification from close-out evaluation or termination gates RFQ / PO / contract / WO',
  'Not native', 'Not native', 'Not native', 'Native — Supplier Qualification rule sets, questionnaires, requalification', 'Native — Ariba SLP qualification projects, disqualification, requalification', 'Ahead', 'Ahead of ERPNext / Odoo / Zoho; comparable to Oracle SQM and Ariba SLP'),
 ('A Onboarding', 'Compliance documents, insurance, expiry gate', 'Compliance Center; vendor Documents tab',
  'Required documents per vendor type, versions, verification queue, insurance cover minimums, expiry reminders, payment gate (Stop / Warn / Off)',
  'Partial — attachments; no expiry gate', 'Partial — documents app', 'Partial — attachments', 'Partial — qualification attachments / certificates', 'Partial — Ariba SLP certificates', 'Ahead', 'Insurance-cover minimums and payment gate are construction-specific'),
 ('A Onboarding', 'Bank account verification', 'Vendor record → Bank tab',
  'Re-enter account no., IFSC, penny-drop style verification status, payments-enabled flag',
  'Partial — bank account record', 'Partial — bank account record', 'Partial — bank details', 'Native — supplier bank accounts with approval', 'Native — business partner bank data', 'On par', ''),
 ('B Sourcing', 'Purchase requisition', 'Purchase Requisitions',
  'Purpose (purchase / manpower / transfer…), items, required by, approval, % ordered / received, Create RFQ',
  'Native — Material Request', 'Partial — Purchase requests via replenishment / Requisition (agreements)', 'Not native in Zoho Books', 'Native — Requisitions', 'Native — Purchase Requisition', 'On par', ''),
 ('B Sourcing', 'RFQ to several vendors + vendor quotes online', 'RFQ & Quotations; public quote link; supplier portal',
  'E-mail compose, secure link with one-time code, line-level quotes, no-bid, CSV upload, questions, T&C, Incoterm',
  'Native — RFQ + supplier portal quotation', 'Native — RFQ / Call for tenders; vendor portal', 'Not native in Zoho Books', 'Native — Sourcing negotiations', 'Native — RFQ / Ariba Sourcing events', 'On par', ''),
 ('B Sourcing', 'Quote comparison & award by line', 'RFQ drawer → comparison, Award by line',
  'Weighted ranking, partial bids, award per line (one PO per vendor), award as contract with BOQ',
  'Partial — Supplier Quotation comparison report', 'Native — compare RFQ lines (call for tenders)', 'Not native', 'Native — award by line', 'Native — award in sourcing', 'On par', 'Award-as-contract carrying the BOQ is construction-specific'),
 ('B Sourcing', 'Blanket order / rate contract', 'Blanket Orders',
  'Agreed lines and rates, call-off POs, allowance %, deadline, consumption tracking',
  'Partial — Blanket Order doctype', 'Native — Purchase agreements (blanket orders)', 'Not native', 'Native — Blanket purchase agreements', 'Native — Outline agreements (contracts / scheduling agreements)', 'On par', ''),
 ('C Commitment', 'PO with approval threshold', 'Purchase Orders',
  'Below threshold issued directly, above it goes to approval; lock confirmed orders; revisions; price-list suggestion',
  'Partial — Workflow on PO', 'Native — PO approval above a minimum amount', 'Native — PO approval', 'Native — approval rules', 'Native — release strategy', 'On par', ''),
 ('C Commitment', 'Subcontract / works contract with BOQ, approval, securities', 'Contracts',
  'BOQ from the tender, approval stages by value, performance BG (add / extend / return / encash), advance, retention %, LD, DLP months',
  'Partial — Subcontracting Order (manufacturing subcontracting)', 'Not native (subcontracting is for manufacturing)', 'Not native', 'Native — Complex work PO with progress payment schedule & retainage', 'Partial — service PO / outline agreement, retention on PO', 'Ahead', 'Bank-guarantee register and works-contract terms are not standard in any of the five'),
 ('C Commitment', 'Work order to contractor + acceptance', 'Work Orders; supplier portal Work orders',
  'WO from contract BOQ, WBS, mobilisation gate, contractor accepts / declines in portal, suspend / short-close',
  'Partial — Subcontracting Order', 'Not native', 'Not native', 'Partial — PO acknowledgement in supplier portal', 'Partial — PO confirmation (Ariba Network)', 'Ahead', ''),
 ('D Execution', 'Goods receipt + quality inspection + returns', 'PO record → Receive goods',
  'Accepted / rejected qty, inspection template & parameters, early / late window, over-receipt Stop / Warn, auto return + debit note',
  'Native — Purchase Receipt + Quality Inspection', 'Native — Receipts (+ Quality app)', 'Native — Purchase receives (Zoho Inventory)', 'Native — Receipts & inspection', 'Native — Goods receipt + QM inspection lot', 'On par', ''),
 ('D Execution', 'Measurement book / service entry / work confirmation', 'Measurement Book; Labour Attendance',
  'Dimensioned entries (N × L × B × D), JMS co-signed by contractor (also in portal), inspection, attendance roll-up to measurement',
  'Not native', 'Not native', 'Not native', 'Native — Work confirmations on progress payment schedules', 'Native — Service entry sheets (lean services) with approval', 'Ahead', 'JMS co-signing and dimensioned measurement are construction-specific'),
 ('D Execution', 'Variations / change orders', 'Contract record → Raise change order',
  'Value and quantity lines, approval, raises WO quantity and contract value',
  'Not native', 'Not native', 'Not native', 'Native — PO change orders', 'Native — PO changes / versions', 'On par', ''),
 ('E Bill & Pay', 'Vendor bill with 3-way match & holds', 'Invoices & Payments',
  'Bill against PO / RA bill / direct; 3-way match with qty / amount tolerance; holds; debit / credit notes; instalments',
  'Native — Purchase Invoice against PO / Receipt', 'Native — Bill control policy + 3-way matching', 'Add-on — BillPay 2 / 3-way reconciliation', 'Native — matching tolerances, invoice holds (Qty Rec, Qty Ord, Price)', 'Native — invoice verification with tolerance keys & payment blocks', 'On par', ''),
 ('E Bill & Pay', 'Progress (RA) bill with retention, advance recovery, TDS, cess', 'RA Bills & Certification; Retention & Deductions',
  'RA bill from signed measurements; retention, mobilisation advance recovery, TDS, labour cess, material recovery, penalty; verify → certify → approve → payable',
  'Not native', 'Not native', 'Not native (retention / progress invoicing are on the sales side)', 'Native — progress payments with retainage', 'Partial — retention on PO / invoice', 'Ahead', 'Indian statutory deductions (TDS, cess) built in'),
 ('E Bill & Pay', 'Payment with TDS', 'Invoices & Payments → Record payment / Select payable',
  'Batch "Select payable", TDS by category & threshold, cheque details, advances, write-off',
  'Native — Payment Entry + Tax Withholding Category', 'Native — payments (TDS via localisation)', 'Native — payments, TDS (India)', 'Native — payments, withholding tax', 'Native — payment run, withholding tax', 'On par', ''),
 ('F Closure', 'Close-out, handover, DLP / warranty', 'Close-out & Handover; DLP & Warranty',
  'Punch list → final inspection → handover certificate → DLP defects → goods warranties & claims',
  'Partial — Warranty Claim (sales side)', 'Not native', 'Not native', 'Partial — work confirmation completion', 'Partial — warranty / QM notifications', 'Ahead', ''),
 ('F Closure', 'Final settlement, retention release, contractor release', 'Final Settlement; Retention & Deductions; Contractor Release',
  'Final account statement agreed / disputed; retention release with approval; no-claim release certificate with closing evaluation',
  'Not native', 'Not native', 'Not native', 'Partial — retainage release on progress payments', 'Partial — release of retention', 'Ahead', ''),
 ('F Closure', 'Termination & blacklist', 'Termination & Final Account',
  'Terminate → encash guarantees → final account → blacklist / hold / no action; open WOs short-closed',
  'Partial — supplier on hold / disabled', 'Partial — vendor archive', 'Partial — inactive vendor', 'Partial — supplier inactivation', 'Partial — block supplier / disqualify', 'Ahead', ''),
 ('Cross-cutting', 'Supplier performance scorecard', 'Vendor Scorecard; Performance & Progress',
  'Formula criteria & weights, standings that warn / prevent RFQ & PO, CAPs, closing evaluation feeds the score',
  'Native — Supplier Scorecard (criteria, variables, periods)', 'Partial — on-time delivery rate', 'Not native', 'Native — supplier performance evaluation', 'Native — supplier evaluation (price, quality, delivery, service…)', 'On par', ''),
 ('Cross-cutting', 'Supplier portal', 'Supplier portal (/supplier)',
  'RFQs, POs, work orders, RA claims, daily attendance, punch list, bills & payments, price list, documents, queries, users',
  'Native — supplier portal (RFQ, PO, invoices)', 'Native — vendor portal', 'Native — vendor portal (accept PO, upload invoices)', 'Native — Supplier Portal (PO, ASN, invoices, negotiations)', 'Native — Ariba Network', 'Ahead', 'Contractor-side features (attendance, RA claims, punch list) are beyond the others'),
 ('Cross-cutting', 'Audit trail', 'Audit Log; each record → activity',
  'Every action with user, time, record link', 'Native — version history', 'Native — chatter', 'Native — history', 'Native — audit', 'Native — change documents', 'On par', ''),
 ('Cross-cutting', 'Roles & segregation of duties', '— (removed on request)',
  'Single-user prototype: anyone can act on any step', 'Native — roles & permissions', 'Native — access rights', 'Native — roles', 'Native — roles, SoD', 'Native — roles, SoD', 'Behind', 'Removed by request; needed before production use'),
 ('Cross-cutting', 'Accounting, GL postings, integrations', '— (prototype)',
  'No general ledger, stock valuation, e-invoicing or bank file', 'Native — full ERP', 'Native — full ERP', 'Native — Books accounting', 'Native — Fusion Financials', 'Native — S/4HANA Finance', 'Behind', 'Expected: this is a front-end prototype for the vendor & contract module'),
 ('Screen UX', 'List screens: filters, columns, export', 'Every list',
  'Right-side Filters panel with chips (multi-select), Reset / Apply, column picker, CSV export, footer count',
  'Native — sidebar filters, report view, export', 'Native — search panel filters / group by, export', 'Native — custom views, export', 'Native — saved searches, export', 'Native — Fiori filter bar, variants, export', 'On par', 'Filters panel modelled on your Project Center'),
 ('Screen UX', 'Record screens', 'Drawers with tabs',
  'Record opens in a side drawer with tabs; links to every related record', 'Form view with connections', 'Form view with smart buttons', 'Detail page with related lists', 'Page with tabs / regions', 'Fiori object page with sections', 'Different', 'Drawer keeps the list visible — closer to Odoo / Fiori object pages than to full-page forms'),
]

SOURCES = [
 ('ERPNext', 'Procurement cycle overview', 'https://docs.frappe.io/erpnext/procurement-cycle-overview.md'),
 ('ERPNext', 'Material Request', 'https://docs.frappe.io/erpnext/material-request.md'),
 ('ERPNext', 'Request for Quotation', 'https://docs.frappe.io/erpnext/request-for-quotation.md'),
 ('ERPNext', 'Supplier quotation through the supplier portal', 'https://docs.frappe.io/erpnext/how-to-create-a-supplier-quotation-through-the-supplier-portal.md'),
 ('ERPNext', 'Purchase Order', 'https://docs.frappe.io/erpnext/purchase-order'),
 ('ERPNext', 'Supplier Scorecard', 'https://docs.frappe.io/erpnext/supplier-scorecard'),
 ('ERPNext', 'Subcontracting reports', 'https://docs.frappe.io/erpnext/subcontracting/reports.md'),
 ('Odoo', 'Bill control policies & 3-way matching', 'https://www.odoo.com/documentation/16.0/it/applications/inventory_and_mrp/purchase/manage_deals/control_bills.html'),
 ('Odoo', 'Purchase agreements (blanket orders, call for tenders) — Odoo 17 book', 'https://www.cybrosys.com/odoo/odoo-books/v17/purchase/purchase-agreements/'),
 ('Odoo', '3-way matching — Odoo 17 book', 'https://cybrosys.com/odoo/odoo-books/v17/purchase/3way-matching'),
 ('Zoho', 'Zoho Books purchase orders', 'https://www.zoho.com/books/purchase-order/'),
 ('Zoho', 'Zoho Books vendor portal', 'https://www.zoho.com/books/help/vendor-portal/'),
 ('Zoho', 'Zoho Books BillPay (2 / 3-way reconciliation, vendor approval)', 'https://www.zoho.com/books/accounting-software-features/billpay/'),
 ('Zoho', 'Zoho Books progress invoicing (sales side)', 'https://www.zoho.com/books/help/quote/progress-invoice.html'),
 ('Oracle', 'Supplier qualification rule sets', 'https://docs.oracle.com/en/cloud/saas/procurement/25c/oapro/supplier-qualification-rule-sets.html'),
 ('Oracle', 'Supplier registration options', 'https://docs.oracle.com/en/cloud/saas/procurement/23d/oapro/Chunk192153297.html'),
 ('Oracle', 'POs for complex services with progress payment schedules & retainage', 'https://docs.oracle.com/en/cloud/saas/procurement/25b/oaprc/how-you-create-purchase-orders-for-complex-services-with-progress.html'),
 ('Oracle', 'Work confirmations for progress payment schedules', 'https://docs.oracle.com/en/cloud/saas/procurement/25d/oaprc/how-you-create-work-confirmations-for-purchase-orders-with-progress-payments-schedule.html'),
 ('Oracle', 'Payables invoice hold types', 'https://docs.oracle.com/en/cloud/saas/financials/25c/fappp/types-of-holds.html'),
 ('Oracle', 'How invoices are validated', 'https://docs.oracle.com/en/cloud/saas/financials/25a/fappp/how-invoices-are-validated.html'),
 ('Oracle', 'Supplier Portal overview (PDF)', 'https://www.oracle.com/assets/supplier-portal-1558356.pdf'),
 ('SAP', 'Purchase order processing in S/4HANA', 'https://learning.sap.com/courses/exploring-operational-procurement-in-sap-s-4hana/outlining-purchase-order-processing-in-sap-s-4hana'),
 ('SAP', 'Processing a service entry sheet', 'https://learning.sap.com/learning-journeys/exploring-the-basics-of-external-procurement-in-sap-s-4hana/processing-a-service-entry-sheet_e42ab0b6-291a-4dad-9f5c-7f508d21ae5c'),
 ('SAP', 'Procurement of services (S/4HANA Cloud)', 'https://learning.sap.com/courses/implementing-sap-s-4hana-cloud-public-edition-sourcing-and-procurement/carrying-out-procurement-of-services-process-steps'),
 ('SAP', 'Purchase orders with retention', 'https://learning.sap.com/courses/describing-consumable-purchasing/creating-purchase-orders-with-retention_d1eb3aa4-60a7-441c-bf64-2e4703a52956'),
 ('SAP', 'Supplier evaluation', 'https://learning.sap.com/courses/purchasing-in-sap-s-4hana/evaluating-suppliers-using-logistics-information-system'),
 ('SAP', 'Ariba Supplier Lifecycle & Performance', 'https://learning.sap.com/courses/discovering-sap-ariba-supplier-management/explaining-sap-ariba-s-supplier-lifecycle-and-performance-processes_cc7c5e73-dc08-4acc-88eb-c9d5bd4e5bf7'),
 ('SAP', 'Supplier qualifications (Ariba)', 'https://learning.sap.com/courses/managing-supplier-qualifications-and-miscellaneous-processes/characterizing-supplier-qualifications'),
]

GAPS = [
 ('Behind', 'Roles, permissions, segregation of duties', 'All five platforms', 'Removed at your request. Needed before real users: who can approve, pay, release retention.'),
 ('Behind', 'Accounting postings, stock valuation, GL', 'All five platforms', 'Prototype keeps records only. Connect to your ERP / Tally for postings.'),
 ('Behind', 'Real e-mail / e-invoice (IRN) / bank payment file', 'Oracle, SAP, Zoho (India), ERPNext (India)', 'E-mails and payments are simulated in the prototype.'),
 ('Behind', 'Reverse auctions / sealed-bid sourcing events', 'Oracle Sourcing, SAP Ariba Sourcing', 'RFQ only; add if competitive bidding events are needed.'),
 ('Behind', 'Contract clause library & document authoring', 'Oracle Procurement Contracts, SAP Ariba Contracts', 'Contracts carry terms as text / templates only.'),
 ('Fixed in this UAT', 'Document numbers missing on PO, RFQ and bill lists', 'All five show the number first', 'Added PO no., RFQ no. and Bill no. columns.'),
 ('Fixed in this UAT', 'PO → RFQ and bill → PO references not clickable', 'Odoo smart buttons, ERPNext connections, SAP document flow', 'Now links that open the source record.'),
 ('Ahead', 'Construction chain: measurement / JMS → RA bill → retention, advance, TDS, cess → DLP → final settlement → release', 'ERPNext, Odoo, Zoho have none; Oracle (complex work) and SAP (SES, retention) cover parts', 'Your main differentiator — keep it.'),
 ('Ahead', 'Contractor portal: attendance, RA claims, JMS agreement, punch list', 'None of the five offer these contractor-side screens', 'Keep.'),
 ('Ahead', 'Bank-guarantee register (add, extend, return, encash)', 'Not standard in the five', 'Keep.'),
]

wb = Workbook(); wb.remove(wb.active)
F = load('final'); M = load('mapping')
fr = F['R']; mr = M['R']
ok = lambda rr: sum(1 for x in rr if x['result'] == 'PASS')
ws = wb.create_sheet('Summary')
lines = [
 ('NebullaOne — Workflow UAT, field mapping and platform comparison', True),
 (f"Workflow UAT (final.js): {ok(fr)}/{len(fr)} lifecycle steps pass — vendor procure-to-pay V01–V14, contractor contract-to-close C01–C18, data chain", False),
 (f"Field mapping (mapping.js): {ok(mr)}/{len(mr)} checks pass — P2P 8, contract-to-close 9, data-wide invariants 7", False),
 ('Two defects found and fixed: document numbers missing on PO / RFQ / bill lists; PO → RFQ and bill → PO references not clickable.', False),
 (f"Platform comparison: {len(CMP)} steps vs ERPNext, Odoo, Zoho, Oracle Fusion, SAP S/4HANA + Ariba — Ahead {sum(1 for c in CMP if c[9]=='Ahead')}, On par {sum(1 for c in CMP if c[9]=='On par')}, Different {sum(1 for c in CMP if c[9]=='Different')}, Behind {sum(1 for c in CMP if c[9]=='Behind')}", False),
 ('Status words: Native = standard feature · Partial = covered in part / by configuration · Add-on = paid add-on or separate app · Not native = not available as standard.', False),
 ('Sources: vendor documentation listed on the Sources sheet (checked October 2026). Odoo items marked "book" cite the Odoo 17 functional book where the official page was not found.', False),
]
for i, (t, b) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = Font(bold=b, size=14 if b else 11)
ws.column_dimensions['A'].width = 150

sheet(wb, 'Workflow UAT', ['Step', 'Scenario', 'Actual result', 'Result'], [[x['id'], x['scn'], x['actual'], x['result']] for x in fr], [8, 60, 90, 9], colour_cols=(3,))
sheet(wb, 'Field mapping', ['Check', 'What maps to what', 'Actual result', 'Result'], [[x['id'], x['scn'], x['actual'], x['result']] for x in mr], [9, 70, 100, 9], colour_cols=(3,))
sheet(wb, 'Platform comparison', ['Stage', 'Workflow step', 'NebullaOne screen', 'NebullaOne has', 'ERPNext', 'Odoo', 'Zoho (Books / Inventory)', 'Oracle Fusion Procurement', 'SAP S/4HANA + Ariba', 'NebullaOne vs others', 'Note'],
      [list(c) for c in CMP], [13, 26, 26, 46, 26, 26, 26, 30, 30, 12, 34], colour_cols=(4, 5, 6, 7, 8, 9), height=75)
sheet(wb, 'Gaps & recommendations', ['Position', 'Topic', 'Who has it', 'Recommendation'], [list(g) for g in GAPS], [16, 52, 46, 70], colour_cols=(0,))
sheet(wb, 'Sources', ['Platform', 'Page', 'URL'], [list(s) for s in SOURCES], [12, 60, 110])
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT); print('wrote', OUT)
