# Builds docs/UAT_Report_NebullaOne.xlsx from out/res-uat.json and the end-to-end suite results.
# Sheets: Summary (per page), Issues (fail / warn with what was done), All checks, End-to-end suites.
import json, os, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'docs', 'UAT_Report_NebullaOne.xlsx')
R = json.load(open(os.path.join(HERE, 'out', 'res-uat.json')))
FIXED = json.load(open(sys.argv[1])) if len(sys.argv) > 1 else []  # issues found in the first run and fixed before this run

HEAD = PatternFill('solid', fgColor='1F3A5F'); HF = Font(bold=True, color='FFFFFF')
FILL = {'PASS': 'E3F4E8', 'FAIL': 'FBE3E3', 'WARN': 'FFF2D6', 'INFO': 'EEF1F5', 'Fixed': 'E3F4E8'}
thin = Side(style='thin', color='D7DCE3'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(wb, title, cols, rows, widths, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]: c.fill = HEAD; c.font = HF; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical='top', wrap_text=True); c.border = BOX
        if status_col is not None:
            v = row[status_col].value
            if v in FILL: row[status_col].fill = PatternFill('solid', fgColor=FILL[v]); row[status_col].font = Font(bold=True)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    return ws

wb = Workbook(); wb.remove(wb.active)
cnt = collections.Counter(r['result'] for r in R)
pages = collections.OrderedDict()
for r in R:
    k = (r['module'], r['page']); pages.setdefault(k, collections.Counter())[r['result']] += 1
summary_rows = []
for (m, pg), c in pages.items():
    st = 'FAIL' if c['FAIL'] else 'WARN' if c['WARN'] else 'PASS'
    summary_rows.append([m, pg, sum(c.values()), c['PASS'], c['FAIL'], c['WARN'], c['INFO'], st])
ws = sheet(wb, 'Summary', ['Module', 'Page / area', 'Checks', 'Pass', 'Fail', 'Warn', 'Info', 'Result'], summary_rows, [20, 32, 9, 9, 9, 9, 9, 11], status_col=7)
ws.insert_rows(1, 4)
ws['A1'] = 'NebullaOne — Vendor & Contract module UAT'; ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = f"{len(R)} checks · {cnt['PASS']} pass · {cnt['FAIL']} fail · {cnt['WARN']} warn · {cnt['INFO']} info · {len(FIXED)} issues found and fixed during UAT"
ws['A3'] = 'Every page of both modules, the supplier portal (contractor and goods supplier) and the public pages; then a 1280 px laptop layout pass.'
ws.freeze_panes = 'A6'; ws.auto_filter.ref = f"A5:H{ws.max_row}"

issues = [[f['page'], f['area'], f['check'], 'Fixed', f['found'], f['fix']] for f in FIXED]
issues += [[r['page'], r['area'], r['check'], r['result'], r['detail'], 'Open — see note'] for r in R if r['result'] in ('FAIL', 'WARN')]
sheet(wb, 'Issues', ['Page', 'Area', 'Check', 'Status', 'What was seen', 'Action'], issues, [26, 11, 50, 10, 55, 60], status_col=3)

sheet(wb, 'All checks', ['Module', 'Page', 'Area', 'Check', 'Result', 'Detail'],
      [[r['module'], r['page'], r['area'], r['check'], r['result'], r['detail']] for r in R], [18, 28, 10, 70, 9, 70], status_col=4)

suites = []
for s, what in [('merge', 'Vendor master: bank verification, foreign vendor, formats, uploads, document versions'), ('fix1', 'Vendor gates and approvals'),
                ('fix23', 'Award → contract, contract approval, guarantees, work-order gates, change orders'), ('fix4', 'Inspection / NCR, quantities, material, equipment, DPR, JMS, WBS'),
                ('fix5', 'Close-out, final bill, retention, portal invoices'), ('final', 'Full vendor + contractor lifecycle (33 steps)'), ('bench', 'Benchmark fields (19)'),
                ('flow', 'Requisition → RFQ → PO → GRN, holds, change orders, audit log'), ('closure', 'Final settlement, DLP & warranty, release, termination, requalification'),
                ('wire', 'Every menu page, record links, deep links, filters panel, public pages')]:
    f = os.path.join(HERE, 'out', f'res-{s}.json')
    if not os.path.exists(f): continue
    d = json.load(open(f)); rr = d['R'] if isinstance(d, dict) else d
    ok = sum(1 for x in rr if x.get('result') == 'PASS'); bad = len(rr) - ok
    suites.append([s + '.js', what, len(rr), ok, bad, 'PASS' if not bad and not (d.get('errs') if isinstance(d, dict) else None) else 'FAIL'])
sheet(wb, 'End-to-end suites', ['Suite', 'Covers', 'Checks', 'Pass', 'Fail', 'Result'], suites, [12, 70, 9, 9, 9, 10], status_col=5)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT); print('wrote', OUT, len(R), 'checks', dict(cnt))
