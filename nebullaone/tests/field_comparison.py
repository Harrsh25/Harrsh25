# Field-by-field comparison: NebullaOne vs ERPNext, Odoo, Zoho Books, Oracle Fusion, SAP S/4HANA
# Y = standard field, A = via a standard add-on / localisation (e.g. ERPNext India Compliance), - = not a standard field
import openpyxl, sys
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
OUT = sys.argv[1] if len(sys.argv) > 1 else 'docs/Field_Comparison_NebullaOne.xlsx'
H = ['Document', 'Field', 'ERPNext', 'Odoo', 'Zoho Books', 'Oracle', 'SAP', 'NebullaOne', 'Where in NebullaOne', 'Note']
HAD, ADDED, REMOVED, NA = 'Had it', 'Added now', 'Removed at your request', 'Not added'
R = []
def add(doc, field, e, o, z, orc, s, st, where='', note=''): R.append([doc, field, e, o, z, orc, s, st, where, note])
V = 'Vendor'
for f, m, w in [
  ('Name / legal name', 'YYYYY', 'Vendor form → Basics'), ('Supplier type / entity type', 'YYYYY', 'Vendor form'), ('Supplier group', 'YYYYY', 'Vendor form'),
  ('Tax ID (GSTIN / PAN / VAT)', 'YYYYY', 'Vendor form → Tax'), ('GST treatment / tax category', 'YAYYY', 'Vendor form → Tax'), ('Place of supply', 'AAY--', 'Vendor form → Tax'),
  ('MSME / Udyam no.', 'AAY--', 'Vendor form → Tax'), ('D-U-N-S number', '---YY', 'Vendor form'), ('Withholding tax (TDS section)', 'AAYYY', 'Vendor form → Tax'),
  ('Currency', 'YYYYY', 'Vendor form'), ('Payment terms', 'YYYYY', 'Vendor form'), ('Payment method', 'YYYYY', 'Vendor form'), ('Default price list', 'YYY--', 'Vendor form'),
  ('Credit limit', '---Y-', 'Vendor form'), ('Contacts (many)', 'YYYYY', 'Vendor → Overview'), ('Addresses / sites (many)', 'YYYYY', 'Vendor → Overview'),
  ('Bank accounts', 'YYYYY', 'Vendor → Bank'), ('Website', 'YYYYY', 'Vendor form'), ('Hold / block with reason', 'YYYYY', 'Vendor → Status & flags'),
  ('Preferred supplier flag', '-YYY-', 'Vendor list star'), ('Buyer / purchase representative', '-Y-YY', 'Vendor form'), ('Purchase warning message', 'YY---', 'Vendor form'),
  ('Receipt reminder', '-Y---', 'Vendor form'), ('Tags', 'YYY--', 'Vendor form'), ('Portal user access', 'YYYYY', 'Vendor portal'), ('Logo / image', 'YYY--', 'Vendor form')]:
  add(V, f, *m, HAD, w)
add(V, 'Company registration no. (CIN / LLPIN)', 'A', '-', '-', 'Y', 'Y', ADDED, 'Vendor form → Tax (validated, 21 chars or LLPIN)', 'Oracle registry ID, SAP commercial register no.')
add(V, 'Print language', 'Y', 'Y', 'Y', '-', 'Y', REMOVED, '', 'You asked to remove it earlier; not re-added')
add(V, 'Default payable account', 'Y', 'Y', 'Y', 'Y', 'Y', REMOVED, '', 'Accounting lives in your finance system; removed earlier at your request')
add(V, 'Opening balance', '-', '-', 'Y', '-', '-', REMOVED, '', 'Removed earlier at your request')
P = 'Purchase order'
for f, m, w in [('Supplier', 'YYYYY', 'New PO'), ('Order date', 'YYYYY', 'Auto'), ('Delivery / expected date (header)', 'YYYYY', 'New PO'),
  ('Items: description, unit, qty, rate', 'YYYYY', 'New PO lines'), ('Currency & exchange rate', 'YYYYY', 'PO details'), ('Payment terms', 'YYYYY', 'PO details'),
  ('Incoterm + named place', 'YYYYY', 'PO details'), ('Ship-to address / warehouse', 'YYYYY', 'PO details'), ('Buyer', 'YYYYY', 'PO details'), ('Project / cost centre', 'YYAYY', 'PO details'),
  ('Additional discount', 'YYYY-', 'PO details'), ('Tax template / category', 'YYYYY', 'PO details'), ('Source document (RFQ / agreement)', 'YYYYY', 'PO header link'),
  ('Vendor quote reference', 'YYYYY', 'PO header'), ('Receipt tolerance', 'YYAYY', 'New PO'), ('Billing policy (ordered / received qty)', 'YYAYY', 'New PO'), ('Notes / remarks', 'YYYYY', 'PO details')]:
  add(P, f, *m, HAD, w)
add(P, 'Line HSN / SAC code', 'A', 'A', 'Y', '-', 'A', ADDED, 'New PO → each line', 'India GST code; validated 4/6/8 digits')
add(P, 'Line tax (GST %)', 'Y', 'Y', 'Y', 'Y', 'Y', ADDED, 'New PO → each line; footer shows GST and total incl. GST', 'Allowed: 0, 5, 12, 18, 28%')
add(P, 'Line need-by / delivery date', 'Y', 'Y', '-', 'Y', 'Y', ADDED, 'New PO → each line; shown under the item in the PO', "Odoo 'Expected Arrival', Oracle 'Need-by date', SAP item delivery date")
add(P, 'Terms & conditions', 'Y', 'Y', 'Y', 'Y', 'Y', ADDED, 'PO details → Terms & printing')
add(P, 'Letter head / print heading', 'Y', '-', 'Y', '-', '-', REMOVED, '', 'Removed earlier at your request')
add(P, 'Drop-ship address', 'Y', 'Y', 'Y', 'Y', 'Y', REMOVED, '', 'Removed earlier at your request')
Q = 'RFQ / quotation'
for f, m, w in [('Title, vendors, due date', 'YYYYY', 'New RFQ'), ('Lines with required-by date', 'YYAYY', 'New RFQ'), ('Terms & conditions', 'YYYYY', 'New RFQ'),
  ('Currency, incoterm, payment terms', 'YYYYY', 'New RFQ / details'), ('Open (preview) date, award date', '---YY', 'RFQ details'), ('Quote validity, lead time, discount, GST per line', 'YYYYY', 'Vendor quote'),
  ('Vendor quote number', 'YYYYY', 'Vendor quote'), ('E-mail message to vendors', 'YYYYY', 'Compose & send')]:
  add(Q, f, *m, HAD, w)
add(Q, 'Agreement type (blanket / call for tender)', '-', 'Y', '-', 'Y', 'Y', REMOVED, '', 'Removed earlier at your request')
G = 'Goods receipt'
for f, m, w in [('Receipt date / posting date', 'YYYYY', 'Receive goods'), ('Received / accepted / rejected qty', 'YYAYY', 'Receive goods'), ('Rejected warehouse', 'YYAYY', 'Receive goods'),
  ('Batch / serial no.', 'YYYYY', 'Receive goods'), ('Quality inspection', 'YYAYY', 'Receive goods'), ('Supplier delivery note / challan', 'YYYYY', 'GRN details'),
  ('Transporter, vehicle no., vehicle date', 'YA--Y', 'GRN details')]:
  add(G, f, *m, HAD, w)
add(G, 'e-Way bill no.', 'A', 'A', 'Y', '-', 'A', ADDED, 'GRN details → Transport (12 digits)', 'ERPNext via India Compliance app')
add(G, 'LR / bill of lading no. and date', 'Y', 'A', '-', 'Y', 'Y', ADDED, 'GRN details → Transport')
add(G, 'Received by', '-', 'Y', '-', 'Y', 'Y', ADDED, 'GRN details → Header (defaults to you)')
B = 'Bill / invoice'
for f, m, w in [('Vendor invoice no. and date', 'YYYYY', 'New bill'), ('Due date / instalments', 'YYYYY', 'New bill'), ('Posting date, supply date', 'YYYYY', 'Bill details'),
  ('Lines against PO with 3-way match', 'YYAYY', 'New bill'), ('GST %', 'YYYYY', 'New bill'), ('Reverse charge', 'AYYYY', 'Bill details → Tax category'),
  ('Hold with reason and release date', 'YYYYY', 'Bill → Hold'), ('Credit / debit notes', 'YYYYY', 'Bill'), ('Currency', 'YYYYY', 'Bill details')]:
  add(B, f, *m, HAD, w)
add(B, 'Place of supply', 'A', 'A', 'Y', '-', 'A', ADDED, 'Bill details → Taxes (defaults from the vendor)')
add(B, 'ITC eligibility', 'A', '-', 'Y', '-', '-', ADDED, 'Bill details → Taxes', 'Zoho Books / ERPNext India values')
PY = 'Payment'
for f, m, w in [('Date, amount, mode, reference', 'YYYYY', 'Record payment'), ('Paid-from and paid-to bank', 'YYYYY', 'Record payment'), ('TDS deducted', 'AAY-A', 'Record payment'),
  ('Payment type (pay / advance / refund)', 'YYYYY', 'Payment details'), ('Value date', 'Y--YY', 'Record payment')]:
  add(PY, f, *m, HAD, w)
add(PY, 'Bank charges', 'Y', 'Y', 'Y', 'Y', 'Y', ADDED, 'Payment details', 'Zoho Books bank charges, ERPNext deductions')
add(PY, 'Cheque / reference date', 'Y', '-', '-', 'Y', 'Y', ADDED, 'Payment details')
C = 'Contract'
for f, m, w in [('Contractor, title, type, value', 'YYYYY', 'Create contract'), ('Start / end', 'YYYYY', 'Create contract'), ('Owner, signatory, signed on', 'Y-YYY', 'Contract'),
  ('Retention, advance & recovery, LD, DLP, BG', '--AYA', 'Create contract', ), ('BOQ lines', '--AYY', 'Create contract'), ('Change orders', '--AYY', 'Contract')]:
  add(C, f, *m, HAD, w)
add(C, 'Payment due (days after certification)', '-', '-', '-', 'Y', 'Y', ADDED, 'Create contract; shown on the contract', 'Default 30 days')
add(C, 'Termination notice (days)', '-', '-', '-', 'Y', 'Y', ADDED, 'Create contract; shown on the contract', 'Default 15 days')

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Fields'
ws.append(H)
fill = {HAD: 'E8F5E9', ADDED: 'E3F2FD', REMOVED: 'FFF3E0', NA: 'FFEBEE'}
for r in R:
  ws.append(r); c = ws.cell(ws.max_row, 8); c.fill = PatternFill('solid', fgColor=fill[r[7]])
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79')
for i, w in enumerate([16, 40, 9, 7, 11, 8, 6, 22, 44, 46], 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
  for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
s = wb.create_sheet('Summary'); s.append(['Document', 'Fields compared', HAD, ADDED, REMOVED])
for d in dict.fromkeys(r[0] for r in R):
  rs = [r for r in R if r[0] == d]; s.append([d, len(rs)] + [sum(r[7] == k for r in rs) for k in (HAD, ADDED, REMOVED)])
s.append(['Total', len(R)] + [sum(r[7] == k for r in R) for k in (HAD, ADDED, REMOVED)])
for c in s[1]: c.font = Font(bold=True)
for i, w in enumerate([18, 16, 10, 12, 24], 1): s.column_dimensions[get_column_letter(i)].width = w
k = wb.create_sheet('Key & sources')
for row in [['Y', 'Standard field in the product'], ['A', 'Through a standard add-on or country localisation (e.g. ERPNext India Compliance, Odoo l10n_in, SAP country version)'], ['-', 'Not a standard field'], [],
  ['Sources'], ['ERPNext purchase order', 'https://docs.frappe.io/erpnext/purchase-order'], ['ERPNext India e-Way bill (add-on)', 'https://ecosire.com/apps/erpnext/erpnext-eway-bill-manager'],
  ['ERPNext ITC / place of supply (GSTR-2B)', 'https://ecosire.com/apps/erpnext/erpnext-gstr2b-reconciliation'], ['Odoo PO line expected arrival', 'https://github.com/odoo/odoo/pull/221380'],
  ['Zoho Books bank charges on payments', 'https://www.zoho.com/ca/books/kb/banking/record-bank-charges.md'], ['Zoho Books purchase orders', 'https://www.zoho.com/books/purchase-order/'],
  ['Oracle PO schedule (need-by date)', 'https://docs.oracle.com/en/cloud/saas/procurement/25d/oedmp/polinelocationsgt-21774.html'],
  ['SAP S/4HANA PO processing', 'https://learning.sap.com/courses/exploring-operational-procurement-in-sap-s-4hana/outlining-purchase-order-processing-in-sap-s-4hana'],
  ['SAP payment terms', 'https://learning.sap.com/courses/customizing-core-settings-in-financial-accounting-in-sap-s4hana/configuring-payment-terms-and-cash-discounts'],
  [], ['Note', 'Marks come from product documentation and standard forms; field names differ by product, so rows compare the business field, not the label.']]:
  k.append(row)
k.column_dimensions['A'].width = 40; k.column_dimensions['B'].width = 110
wb.save(OUT)
print(OUT, len(R), 'fields;', {x: sum(r[7] == x for r in R) for x in (HAD, ADDED, REMOVED)})
