import csv
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

CSV_PATH = "/root/.claude/uploads/88ea236d-50b0-5ace-9b45-3e67282b94dc/8478b0d3-LILO_OF_220kV_SC_DAUSA__ALWAR.csv"
OUT_PATH = "/home/user/Harrsh25/LILO_220kV_SC_DAUSA_ALWAR_WBS.xlsx"

HEADERS = ["Activity_title", "Activity_Description", "Task_Name", "Task_Description", "Sub_Task_Name"]

T = Side(style="thin", color="000000")
M = Side(style="medium", color="000000")

def bdr(thick=False):
    s = M if thick else T
    return Border(left=s, right=s, top=s, bottom=s)

def fnt(bold=False, sz=10):
    return Font(bold=bold, size=sz, name="Calibri", color="000000")

def aln(h="left", wrap=True):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LILO 220kV WBS"

# Header row
for col_idx, hdr in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col_idx, value=hdr)
    cell.font = fnt(bold=True, sz=10)
    cell.alignment = aln(h="center", wrap=True)
    cell.border = bdr(thick=True)
    cell.fill = PatternFill(fill_type=None)

# Read CSV (skip header row in file)
with open(CSV_PATH, newline='', encoding='utf-8') as f:
    reader = csv.reader(f)
    next(reader)  # skip CSV header
    for row_idx, row in enumerate(reader, 2):
        while len(row) < 5:
            row.append("")

        # Determine level for styling
        if row[0].strip():
            bold, sz = True, 11   # Activity level
        elif row[2].strip():
            bold, sz = True, 10   # Task level
        elif row[4].strip():
            bold, sz = False, 9   # Sub-task level
        else:
            bold, sz = False, 9

        for col_idx, val in enumerate(row[:5], 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val.strip() if val else "")
            cell.font = fnt(bold=bold, sz=sz)
            cell.alignment = aln()
            cell.border = bdr()
            cell.fill = PatternFill(fill_type=None)

# Column widths
for i, w in enumerate([30, 30, 30, 30, 35], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"

wb.save(OUT_PATH)
print(f"Saved: {OUT_PATH}, Rows: {ws.max_row}")
