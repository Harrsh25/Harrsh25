import csv
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

CSV_PATH = "/root/.claude/uploads/88ea236d-50b0-5ace-9b45-3e67282b94dc/8478b0d3-LILO_OF_220kV_SC_DAUSA__ALWAR.csv"
OUT_PATH = "/home/user/Harrsh25/LILO_220kV_SC_DAUSA_ALWAR_WBS.xlsx"

HEADERS = [
    "Level1_Name","Level1_Desc",
    "Level2_Name","Level2_Desc",
    "Level3_Name","Level3_Desc",
    "Level4_Name","Level4_Desc",
    "Level5_Name","Level5_Desc",
    "Assign_To","Timeline","Weight","Tower_Tagging","Leaf_Node_Progress_Type"
]

T = Side(style="thin", color="000000")
M = Side(style="medium", color="000000")

def bdr(thick=False):
    s = M if thick else T
    return Border(left=s, right=s, top=s, bottom=s)

def fnt(bold=False, sz=10):
    return Font(bold=bold, size=sz, name="Calibri", color="000000")

def aln(h="left", wrap=True):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "LILO 220kV WBS"

# Header row
for col_idx, hdr in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col_idx, value=hdr)
    cell.font = fnt(bold=True, sz=10)
    cell.alignment = aln(h="center", wrap=True)
    cell.border = bdr(thick=True)
    cell.fill = PatternFill(fill_type=None)

# Map LILO CSV columns to 15-column format:
# Activity_title     -> Level1_Name (col A / index 0)
# Activity_Description -> Level1_Desc (col B / index 1)
# Task_Name          -> Level3_Name (col E / index 4)  [skipping Level2]
# Task_Description   -> Level3_Desc (col F / index 5)
# Sub_Task_Name      -> Level5_Name (col I / index 8)
# Leaf_Node_Progress_Type = "Binary" when Sub_Task_Name is filled, else blank

with open(CSV_PATH, newline='', encoding='utf-8') as f:
    reader = csv.reader(f)
    next(reader)  # skip CSV header
    for row_idx, src in enumerate(reader, 2):
        while len(src) < 5:
            src.append("")

        act_title = src[0].strip()
        act_desc  = src[1].strip()
        task_name = src[2].strip()
        task_desc = src[3].strip()
        subtask   = src[4].strip()

        # Build 15-column row
        row = [""] * 15
        row[0]  = act_title   # Level1_Name
        row[1]  = act_desc    # Level1_Desc
        row[4]  = task_name   # Level3_Name
        row[5]  = task_desc   # Level3_Desc
        row[8]  = subtask     # Level5_Name
        row[10] = ""          # Assign_To
        row[11] = ""          # Timeline
        row[12] = ""          # Weight
        row[13] = ""          # Tower_Tagging
        row[14] = "Binary" if subtask else ""  # Leaf_Node_Progress_Type

        # Styling by level
        if act_title:
            bold, sz = True, 11
        elif task_name:
            bold, sz = True, 10
        else:
            bold, sz = False, 9

        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = fnt(bold=bold, sz=sz)
            cell.alignment = aln()
            cell.border = bdr()
            cell.fill = PatternFill(fill_type=None)

# Column widths matching 132 KV format
col_widths = [25, 30, 25, 30, 25, 30, 25, 30, 25, 30, 18, 25, 8, 14, 22]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
wb.save(OUT_PATH)
print(f"Saved: {OUT_PATH}, Rows: {ws.max_row}")
